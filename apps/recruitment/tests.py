import io
from datetime import date
from io import StringIO
from unittest import mock
from xml.dom import minidom
from zipfile import ZipFile

import openpyxl
from django.contrib.auth import get_user_model
from django.contrib.messages import get_messages
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.db import IntegrityError, transaction
from django.test import TestCase
from django.urls import reverse
from rest_framework import status
from rest_framework.test import APITestCase

from apps.employment.models import Contrato, Empleado
from apps.locations.models import Area
from apps.organizations.models import OrganizationalLevel, OrganizationNode
from apps.persons.models import Persona
from apps.positions.models import EstatusPosicion, Posicion, Puesto, TipoRequisicion
from apps.recruitment import exports_word
from apps.recruitment.exports import PLANTILLAS_DIR, generar_excel
from apps.recruitment.exports_word import PLANTILLA as PLANTILLA_WORD, _ancestro, _texto_de, generar_word
from apps.recruitment.models import (
    AprobacionRequisicion,
    CompetenciaConductual,
    ConformidadDescriptivo,
    DescriptivoPuesto,
    DiasPorLaborar,
    EstadoRequisicion,
    EtapaAprobacion,
    FuncionPuesto,
    HorarioACubrir,
    IndicadorDesempeno,
    RangoEdad,
    RecursoAsignado,
    Requisicion,
    RolConformidad,
    TipoContratoOfrecido,
)
from apps.recruitment.services import copiar_version, crear_borrador
from apps.users.models import UserRole


class RecruitmentTestDataMixin:
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        call_command("seed_tenant", stdout=StringIO())
        call_command("seed_organizational_levels", stdout=StringIO())
        call_command("seed_user_roles", stdout=StringIO())
        # TipoRequisicion no tiene seed propio (en producción lo llena el
        # import de la sábana) -- se crea a mano, igual que EstatusPosicion
        # abajo, antes de que seed_recruitment_catalogs marque cuál lo
        # requiere.
        cls.tipo_nueva = TipoRequisicion.objects.create(name="Nueva Posición")
        cls.tipo_reemplazo = TipoRequisicion.objects.create(name="Reemplazo")
        call_command("seed_recruitment_catalogs", stdout=StringIO())

        empresa_level = OrganizationalLevel.objects.get(code="empresa")
        cls.company_node = OrganizationNode.objects.create(
            level=empresa_level, code="GPA-REC-TEST", name="Empresa de prueba",
        )
        cls.estatus_vacante = EstatusPosicion.objects.create(name="Vacante Activa")
        cls.tipo_nueva.refresh_from_db()
        cls.tipo_reemplazo.refresh_from_db()
        cls.estado_borrador = EstadoRequisicion.objects.get(name="Borrador")
        cls.estado_cubierta = EstadoRequisicion.objects.get(name="Cubierta")

    @classmethod
    def create_posicion(cls, **kwargs):
        posicion = Posicion(
            organization_node=kwargs.pop("organization_node", cls.company_node),
            estatus=kwargs.pop("estatus", cls.estatus_vacante),
            **kwargs,
        )
        posicion.full_clean()
        posicion.save()
        return posicion

    @classmethod
    def create_requisicion(cls, *, posicion=None, tipo=None, estado=None, fecha_solicitud=date(2026, 1, 1), **kwargs):
        requisicion = Requisicion(
            posicion=posicion or cls.create_posicion(),
            tipo=tipo or cls.tipo_reemplazo,
            estado=estado or cls.estado_borrador,
            fecha_solicitud=fecha_solicitud,
            **kwargs,
        )
        requisicion.full_clean()
        requisicion.save()
        return requisicion


class SeedRecruitmentCatalogsTests(TestCase):
    def test_seeds_everything_confirmado_and_is_idempotent(self):
        call_command("seed_tenant", stdout=StringIO())
        call_command("seed_organizational_levels", stdout=StringIO())

        out = StringIO()
        call_command("seed_recruitment_catalogs", stdout=out)
        self.assertEqual(EstadoRequisicion.objects.count(), 8)
        self.assertEqual(EtapaAprobacion.objects.count(), 4)
        self.assertEqual(TipoContratoOfrecido.objects.count(), 2)
        self.assertEqual(HorarioACubrir.objects.count(), 5)
        # Descriptivo de Puesto: valores copiados del Word real.
        self.assertEqual(RangoEdad.objects.count(), 5)
        self.assertEqual(DiasPorLaborar.objects.count(), 3)
        self.assertEqual(CompetenciaConductual.objects.count(), 12)
        self.assertEqual(RecursoAsignado.objects.count(), 10)
        self.assertEqual(RolConformidad.objects.count(), 3)
        self.assertTrue(RolConformidad.objects.get(name="Colaborador").requiere_persona)
        self.assertFalse(RolConformidad.objects.get(name="Capital Humano").requiere_persona)

        out2 = StringIO()
        call_command("seed_recruitment_catalogs", stdout=out2)
        self.assertIn("0 valores nuevos", out2.getvalue())

    def test_marks_nueva_posicion_as_requiring_justificacion(self):
        call_command("seed_tenant", stdout=StringIO())
        call_command("seed_organizational_levels", stdout=StringIO())
        TipoRequisicion.objects.create(name="Nueva Posición")
        TipoRequisicion.objects.create(name="Reemplazo")

        call_command("seed_recruitment_catalogs", stdout=StringIO())

        self.assertTrue(TipoRequisicion.objects.get(name="Nueva Posición").requiere_justificacion)
        self.assertFalse(TipoRequisicion.objects.get(name="Reemplazo").requiere_justificacion)


class RequisicionJustificacionTests(RecruitmentTestDataMixin, TestCase):
    def test_nueva_posicion_requires_justificacion(self):
        requisicion = Requisicion(
            posicion=self.create_posicion(), tipo=self.tipo_nueva, estado=self.estado_borrador,
            fecha_solicitud=date(2026, 1, 1),
        )
        with self.assertRaises(ValidationError) as context:
            requisicion.full_clean()
        self.assertIn("justificacion", context.exception.message_dict)

    def test_reemplazo_does_not_require_justificacion(self):
        requisicion = Requisicion(
            posicion=self.create_posicion(), tipo=self.tipo_reemplazo, estado=self.estado_borrador,
            fecha_solicitud=date(2026, 1, 1),
        )
        requisicion.full_clean()
        requisicion.save()
        self.assertEqual(requisicion.justificacion, "")

    def test_nueva_posicion_passes_with_justificacion(self):
        requisicion = Requisicion(
            posicion=self.create_posicion(), tipo=self.tipo_nueva, estado=self.estado_borrador,
            fecha_solicitud=date(2026, 1, 1), justificacion="Crecimiento de la línea de producción.",
        )
        requisicion.full_clean()
        requisicion.save()


class RequisicionUnaAbiertaPorPosicionTests(RecruitmentTestDataMixin, TestCase):
    def test_no_puede_haber_dos_requisiciones_abiertas_para_la_misma_posicion(self):
        posicion = self.create_posicion()
        self.create_requisicion(posicion=posicion, estado=self.estado_borrador)

        segunda = Requisicion(
            posicion=posicion, tipo=self.tipo_reemplazo, estado=self.estado_borrador,
            fecha_solicitud=date(2026, 2, 1),
        )
        with self.assertRaises(ValidationError) as context:
            segunda.full_clean()
        self.assertIn("posicion", context.exception.message_dict)

    def test_si_puede_haber_otra_una_vez_que_la_anterior_ya_es_terminal(self):
        posicion = self.create_posicion()
        self.create_requisicion(posicion=posicion, estado=self.estado_cubierta)

        segunda = Requisicion(
            posicion=posicion, tipo=self.tipo_reemplazo, estado=self.estado_borrador,
            fecha_solicitud=date(2026, 2, 1),
        )
        segunda.full_clean()
        segunda.save()
        self.assertEqual(Requisicion.objects.filter(posicion=posicion).count(), 2)


class RequisicionSoftDeleteTests(RecruitmentTestDataMixin, TestCase):
    def test_delete_does_not_remove_the_row(self):
        requisicion = self.create_requisicion()

        requisicion.delete()

        self.assertFalse(Requisicion.objects.filter(pk=requisicion.pk).exists())
        self.assertTrue(Requisicion.all_objects.get(pk=requisicion.pk).is_deleted)

    def test_a_soft_deleted_requisicion_no_longer_blocks_a_new_one(self):
        posicion = self.create_posicion()
        requisicion = self.create_requisicion(posicion=posicion, estado=self.estado_borrador)
        requisicion.delete()

        nueva = Requisicion(
            posicion=posicion, tipo=self.tipo_reemplazo, estado=self.estado_borrador,
            fecha_solicitud=date(2026, 2, 1),
        )
        nueva.full_clean()
        nueva.save()


class AprobacionRequisicionTests(RecruitmentTestDataMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.jefe_inmediato = EtapaAprobacion.objects.get(name="Jefe Inmediato")

    def test_requires_quien_aprobo_si_hay_fecha(self):
        aprobacion = AprobacionRequisicion(
            requisicion=self.create_requisicion(), etapa=self.jefe_inmediato, fecha=date(2026, 1, 2),
        )
        with self.assertRaises(ValidationError) as context:
            aprobacion.full_clean()
        self.assertIn("nombre_manual", context.exception.message_dict)

    def test_ok_con_nombre_capturado_a_mano(self):
        aprobacion = AprobacionRequisicion(
            requisicion=self.create_requisicion(), etapa=self.jefe_inmediato,
            fecha=date(2026, 1, 2), nombre_manual="Juan Pérez",
        )
        aprobacion.full_clean()
        aprobacion.save()

    def test_pendiente_sin_fecha_no_exige_nada(self):
        aprobacion = AprobacionRequisicion(requisicion=self.create_requisicion(), etapa=self.jefe_inmediato)
        aprobacion.full_clean()
        aprobacion.save()

    def test_no_puede_haber_dos_filas_para_la_misma_etapa(self):
        requisicion = self.create_requisicion()
        AprobacionRequisicion.objects.create(requisicion=requisicion, etapa=self.jefe_inmediato)

        with self.assertRaises(IntegrityError), transaction.atomic():
            AprobacionRequisicion.objects.create(requisicion=requisicion, etapa=self.jefe_inmediato)

    def test_borrar_una_aprobacion_libera_su_etapa_para_una_nueva(self):
        requisicion = self.create_requisicion()
        original = AprobacionRequisicion.objects.create(requisicion=requisicion, etapa=self.jefe_inmediato)
        original.delete()

        AprobacionRequisicion.objects.create(requisicion=requisicion, etapa=self.jefe_inmediato)


class RecruitmentAPIAuthenticationTests(RecruitmentTestDataMixin, APITestCase):
    def test_all_recruitment_endpoints_require_authentication(self):
        protected_urls = [
            reverse("estadorequisicion-list"),
            reverse("etapaaprobacion-list"),
            reverse("tipocontratoofrecido-list"),
            reverse("horarioacubrir-list"),
            reverse("requisicion-list"),
            reverse("aprobacionrequisicion-list"),
            reverse("rangoedad-list"),
            reverse("diasporlaborar-list"),
            reverse("competenciaconductual-list"),
            reverse("recursoasignado-list"),
            reverse("rolconformidad-list"),
            reverse("descriptivopuesto-list"),
            reverse("descriptivopuesto-vigente"),
            reverse("conformidaddescriptivo-list"),
        ]
        for url in protected_urls:
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)


class RecruitmentCatalogAPITests(RecruitmentTestDataMixin, APITestCase):
    def setUp(self):
        super().setUp()
        self.user = get_user_model().objects.create_user(
            username="recruitment-catalog-user", email="recruitment-catalog@example.com",
            password="strong-test-password",
        )
        self.client.force_authenticate(user=self.user)

    def test_estado_requisicion_is_read_only_for_any_authenticated_user(self):
        list_url = reverse("estadorequisicion-list")
        response = self.client.get(list_url)
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["count"], 8)

        create_response = self.client.post(list_url, {"name": "Nuevo"}, format="json")
        self.assertEqual(create_response.status_code, status.HTTP_405_METHOD_NOT_ALLOWED)


class RequisicionRoleAPITests(RecruitmentTestDataMixin, APITestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        user_model = get_user_model()
        cls.gestor = user_model.objects.create_user(
            username="requisicion-gestor", email="requisicion-gestor@example.com",
            password="strong-test-password", role=UserRole.objects.get(code="capital-humano"),
        )
        cls.gerente = user_model.objects.create_user(
            username="requisicion-gerente", email="requisicion-gerente@example.com",
            password="strong-test-password", role=UserRole.objects.get(code="colaborador"),
        )
        cls.otro_colaborador = user_model.objects.create_user(
            username="requisicion-otro", email="requisicion-otro@example.com",
            password="strong-test-password", role=UserRole.objects.get(code="colaborador"),
        )

    def _payload(self, posicion):
        return {
            "posicion": str(posicion.pk),
            "tipo": str(self.tipo_reemplazo.pk),
            "estado": str(self.estado_borrador.pk),
            "fecha_solicitud": "2026-01-01",
        }

    def test_un_colaborador_cualquiera_puede_crear_una_requisicion(self):
        self.client.force_authenticate(user=self.gerente)
        response = self.client.post(reverse("requisicion-list"), self._payload(self.create_posicion()), format="json")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(Requisicion.objects.get(pk=response.data["id"]).created_by, self.gerente)

    def test_un_colaborador_solo_ve_las_requisiciones_que_el_creo(self):
        propia = Requisicion.objects.create(
            posicion=self.create_posicion(), tipo=self.tipo_reemplazo, estado=self.estado_borrador,
            fecha_solicitud=date(2026, 1, 1), created_by=self.gerente,
        )
        ajena = Requisicion.objects.create(
            posicion=self.create_posicion(), tipo=self.tipo_reemplazo, estado=self.estado_borrador,
            fecha_solicitud=date(2026, 1, 1), created_by=self.otro_colaborador,
        )

        self.client.force_authenticate(user=self.gerente)
        list_response = self.client.get(reverse("requisicion-list"))
        self.assertEqual(list_response.data["count"], 1)
        self.assertEqual(list_response.data["results"][0]["id"], str(propia.pk))

        detail_own = self.client.get(reverse("requisicion-detail", args=[propia.pk]))
        self.assertEqual(detail_own.status_code, status.HTTP_200_OK)

        detail_other = self.client.get(reverse("requisicion-detail", args=[ajena.pk]))
        self.assertEqual(detail_other.status_code, status.HTTP_404_NOT_FOUND)

    def test_un_colaborador_puede_editar_su_propia_requisicion(self):
        propia = Requisicion.objects.create(
            posicion=self.create_posicion(), tipo=self.tipo_reemplazo, estado=self.estado_borrador,
            fecha_solicitud=date(2026, 1, 1), created_by=self.gerente,
        )
        self.client.force_authenticate(user=self.gerente)
        response = self.client.patch(
            reverse("requisicion-detail", args=[propia.pk]), {"area_solicitante": "Mantenimiento"}, format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        propia.refresh_from_db()
        self.assertEqual(propia.area_solicitante, "Mantenimiento")

    def test_un_colaborador_no_puede_borrar_ni_su_propia_requisicion(self):
        propia = Requisicion.objects.create(
            posicion=self.create_posicion(), tipo=self.tipo_reemplazo, estado=self.estado_borrador,
            fecha_solicitud=date(2026, 1, 1), created_by=self.gerente,
        )
        self.client.force_authenticate(user=self.gerente)
        response = self.client.delete(reverse("requisicion-detail", args=[propia.pk]))
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertTrue(Requisicion.objects.filter(pk=propia.pk).exists())

    def test_gestor_ve_y_administra_todas_las_requisiciones(self):
        Requisicion.objects.create(
            posicion=self.create_posicion(), tipo=self.tipo_reemplazo, estado=self.estado_borrador,
            fecha_solicitud=date(2026, 1, 1), created_by=self.gerente,
        )
        ajena = Requisicion.objects.create(
            posicion=self.create_posicion(), tipo=self.tipo_reemplazo, estado=self.estado_borrador,
            fecha_solicitud=date(2026, 1, 1), created_by=self.otro_colaborador,
        )

        self.client.force_authenticate(user=self.gestor)
        list_response = self.client.get(reverse("requisicion-list"))
        self.assertEqual(list_response.data["count"], 2)

        delete_response = self.client.delete(reverse("requisicion-detail", args=[ajena.pk]))
        self.assertEqual(delete_response.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(Requisicion.objects.filter(pk=ajena.pk).exists())

    def _crear(self, usuario, puesto="Operador", **extra):
        puesto_obj = Puesto.objects.get_or_create(name=puesto)[0]
        return Requisicion.objects.create(
            posicion=self.create_posicion(puesto=puesto_obj), tipo=self.tipo_reemplazo, estado=self.estado_borrador,
            fecha_solicitud=extra.pop("fecha_solicitud", date(2026, 1, 1)), created_by=usuario, **extra,
        )

    def test_la_requisicion_trae_la_etiqueta_de_su_posicion_y_quien_la_solicito(self):
        propia = self._crear(self.gerente, puesto="Operador de Soldadura")
        self.gerente.first_name, self.gerente.last_name = "Ana", "Ruiz"
        self.gerente.save()
        self.client.force_authenticate(user=self.gestor)
        datos = self.client.get(reverse("requisicion-detail", args=[propia.pk])).data
        self.assertEqual(datos["posicion_etiqueta"], "Operador de Soldadura — Empresa de prueba")
        self.assertEqual(datos["solicitante"], "Ana Ruiz")
        self.assertEqual(datos["creado_por"], self.gerente.pk)

    def test_una_requisicion_importada_no_tiene_solicitante(self):
        importada = self._crear(None)
        self.client.force_authenticate(user=self.gestor)
        datos = self.client.get(reverse("requisicion-detail", args=[importada.pk])).data
        self.assertIsNone(datos["solicitante"])
        self.assertIsNone(datos["creado_por"])

    def test_la_lista_busca_sin_acentos_y_ordena_solo_por_campos_declarados(self):
        self._crear(self.gestor, puesto="Ingeniero de Servicio Técnico", fecha_solicitud=date(2026, 3, 1))
        self._crear(self.gestor, puesto="Operador de Soldadura", fecha_solicitud=date(2026, 5, 1))
        self.client.force_authenticate(user=self.gestor)
        url = reverse("requisicion-list")

        encontradas = self.client.get(url, {"search": "tecnico"}).data["results"]
        self.assertEqual([r["posicion_etiqueta"].split(" — ")[0] for r in encontradas], ["Ingeniero de Servicio Técnico"])

        por_fecha = self.client.get(url, {"ordering": "fecha_solicitud"}).data["results"]
        self.assertEqual([r["fecha_solicitud"] for r in por_fecha], ["2026-03-01", "2026-05-01"])
        por_fecha_desc = self.client.get(url, {"ordering": "-fecha_solicitud"}).data["results"]
        self.assertEqual([r["fecha_solicitud"] for r in por_fecha_desc], ["2026-05-01", "2026-03-01"])
        # Un campo que la tabla no muestra no se puede usar para ordenar.
        self.assertEqual(self.client.get(url, {"ordering": "sueldo_mensual_neto"}).status_code, status.HTTP_200_OK)

    def test_la_busqueda_de_un_colaborador_no_encuentra_las_de_otros(self):
        self._crear(self.gerente, puesto="Soldador propio")
        self._crear(self.otro_colaborador, puesto="Soldador ajeno")
        self.client.force_authenticate(user=self.gerente)
        etiquetas = [
            r["posicion_etiqueta"] for r in self.client.get(reverse("requisicion-list"), {"search": "soldador"}).data["results"]
        ]
        self.assertEqual(len(etiquetas), 1)
        self.assertIn("Soldador propio", etiquetas[0])

    def test_la_lista_filtra_por_estado_y_tipo(self):
        self._crear(self.gestor, puesto="Uno")
        cubierta = self._crear(self.gestor, puesto="Dos")
        Requisicion.objects.filter(pk=cubierta.pk).update(estado=self.estado_cubierta)
        self.client.force_authenticate(user=self.gestor)
        url = reverse("requisicion-list")
        self.assertEqual(self.client.get(url, {"estado": self.estado_cubierta.pk}).data["count"], 1)
        self.assertEqual(self.client.get(url, {"tipo": self.tipo_nueva.pk}).data["count"], 0)

    def test_el_solicitante_no_se_puede_autorizar_a_si_mismo(self):
        autorizada = EstadoRequisicion.objects.get(code="autorizada")
        pendiente = EstadoRequisicion.objects.get(code="pendiente-de-autorizacion")
        self.client.force_authenticate(user=self.gerente)
        posicion = self.create_posicion()

        # Al crear: no puede elegir un estado que no sea Borrador/Pendiente...
        rechazada = self.client.post(
            reverse("requisicion-list"), {**self._payload(posicion), "estado": str(autorizada.pk)}, format="json",
        )
        self.assertEqual(rechazada.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("estado", rechazada.data)

        # ...pero sí puede crearla, y mandarla a autorización.
        creada = self.client.post(reverse("requisicion-list"), self._payload(posicion), format="json")
        self.assertEqual(creada.status_code, status.HTTP_201_CREATED)
        url = reverse("requisicion-detail", args=[creada.data["id"]])
        enviada = self.client.patch(url, {"estado": str(pendiente.pk)}, format="json")
        self.assertEqual(enviada.status_code, status.HTTP_200_OK)

        # Al editar: tampoco puede saltar a Autorizada.
        salto = self.client.patch(url, {"estado": str(autorizada.pk)}, format="json")
        self.assertEqual(salto.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(Requisicion.objects.get(pk=creada.data["id"]).estado, pendiente)

        # Capital Humano sí.
        self.client.force_authenticate(user=self.gestor)
        self.assertEqual(self.client.patch(url, {"estado": str(autorizada.pk)}, format="json").status_code, 200)

    def test_sin_estado_la_requisicion_empieza_en_borrador(self):
        self.client.force_authenticate(user=self.gerente)
        payload = self._payload(self.create_posicion())
        del payload["estado"]
        respuesta = self.client.post(reverse("requisicion-list"), payload, format="json")
        self.assertEqual(respuesta.status_code, status.HTTP_201_CREATED, respuesta.data)
        self.assertEqual(respuesta.data["estado"], self.estado_borrador.pk)

    def test_editar_otro_campo_no_toca_el_estado_aunque_ya_no_pueda_elegirlo(self):
        propia = self._crear(self.gerente)
        Requisicion.objects.filter(pk=propia.pk).update(estado=self.estado_cubierta)
        self.client.force_authenticate(user=self.gerente)
        respuesta = self.client.patch(
            reverse("requisicion-detail", args=[propia.pk]), {"area_solicitante": "Mantenimiento"}, format="json",
        )
        self.assertEqual(respuesta.status_code, status.HTTP_200_OK)

    def test_el_dueno_puede_exportar_su_propia_requisicion_a_excel(self):
        propia = Requisicion.objects.create(
            posicion=self.create_posicion(), tipo=self.tipo_reemplazo, estado=self.estado_borrador,
            fecha_solicitud=date(2026, 1, 1), created_by=self.gerente,
        )
        self.client.force_authenticate(user=self.gerente)
        response = self.client.get(reverse("requisicion-exportar-excel", args=[propia.pk]))
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_no_se_puede_exportar_la_requisicion_de_otro_colaborador(self):
        ajena = Requisicion.objects.create(
            posicion=self.create_posicion(), tipo=self.tipo_reemplazo, estado=self.estado_borrador,
            fecha_solicitud=date(2026, 1, 1), created_by=self.otro_colaborador,
        )
        self.client.force_authenticate(user=self.gerente)
        response = self.client.get(reverse("requisicion-exportar-excel", args=[ajena.pk]))
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)


class AprobacionRequisicionAPIRoleTests(RecruitmentTestDataMixin, APITestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.jefe_inmediato = EtapaAprobacion.objects.get(name="Jefe Inmediato")
        cls.gestor = get_user_model().objects.create_user(
            username="aprobacion-gestor", email="aprobacion-gestor@example.com",
            password="strong-test-password", role=UserRole.objects.get(code="capital-humano"),
        )
        cls.colaborador = get_user_model().objects.create_user(
            username="aprobacion-colaborador", email="aprobacion-colaborador@example.com",
            password="strong-test-password", role=UserRole.objects.get(code="colaborador"),
        )

    def test_un_colaborador_no_puede_crear_ni_leer_aprobaciones(self):
        self.client.force_authenticate(user=self.colaborador)
        response = self.client.get(reverse("aprobacionrequisicion-list"))
        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)

    def test_gestor_puede_registrar_una_aprobacion(self):
        requisicion = self.create_requisicion()
        self.client.force_authenticate(user=self.gestor)
        response = self.client.post(
            reverse("aprobacionrequisicion-list"),
            {
                "requisicion": str(requisicion.pk), "etapa": str(self.jefe_inmediato.pk),
                "fecha": "2026-01-02", "nombre_manual": "Juan Pérez",
            },
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(AprobacionRequisicion.objects.get(pk=response.data["id"]).created_by, self.gestor)


class BackfillRequisicionesVacantesTests(RecruitmentTestDataMixin, TestCase):
    """
    apps.recruitment.management.commands.backfill_requisiciones_vacantes --
    confirmado 2026-10-01: a las Posicion ya vacantes no se les inventa
    justificación, fecha ni tipo -- lo que falta se reporta, no se adivina.
    """

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.estatus_pendiente = EstatusPosicion.objects.create(name="Vacante Pendiente de Confirmación")
        cls.estatus_suspendida = EstatusPosicion.objects.create(name="Vacante Suspendida")
        cls.estatus_eliminada = EstatusPosicion.objects.create(name="Vacante Eliminada")

    def test_migra_una_vacante_con_tipo_conocido_y_sin_justificacion_requerida(self):
        posicion = self.create_posicion(
            estatus=self.estatus_vacante, tipo_requisicion=self.tipo_reemplazo,
            fecha_registro_vacante=date(2025, 6, 1),
        )
        out = StringIO()
        call_command("backfill_requisiciones_vacantes", stdout=out)

        requisicion = Requisicion.objects.get(posicion=posicion)
        self.assertEqual(requisicion.tipo, self.tipo_reemplazo)
        self.assertEqual(requisicion.estado.name, "En Reclutamiento")
        self.assertEqual(requisicion.fecha_solicitud, date(2025, 6, 1))
        self.assertIsNone(requisicion.created_by)
        self.assertIn("Requisición(es) migrada(s): 1", out.getvalue())

    def test_mapea_cada_estatus_vacante_al_estado_de_requisicion_correcto(self):
        casos = [
            (self.estatus_vacante, "En Reclutamiento"),
            (self.estatus_pendiente, "Pendiente de Autorización"),
            (self.estatus_suspendida, "Suspendida"),
            (self.estatus_eliminada, "Cancelada"),
        ]
        for estatus, estado_esperado in casos:
            with self.subTest(estatus=estatus.name):
                posicion = self.create_posicion(
                    estatus=estatus, tipo_requisicion=self.tipo_reemplazo,
                    fecha_registro_vacante=date(2025, 6, 1),
                )
                call_command("backfill_requisiciones_vacantes", stdout=StringIO())
                self.assertEqual(Requisicion.objects.get(posicion=posicion).estado.name, estado_esperado)

    def test_omite_una_vacante_sin_tipo_requisicion(self):
        posicion = self.create_posicion(estatus=self.estatus_vacante, fecha_registro_vacante=date(2025, 6, 1))
        out = StringIO()
        call_command("backfill_requisiciones_vacantes", stdout=out)

        self.assertFalse(Requisicion.objects.filter(posicion=posicion).exists())
        self.assertIn("sin Tipo de requisición capturado", out.getvalue())

    def test_omite_una_vacante_de_tipo_nueva_posicion_porque_falta_la_justificacion(self):
        posicion = self.create_posicion(
            estatus=self.estatus_vacante, tipo_requisicion=self.tipo_nueva,
            fecha_registro_vacante=date(2025, 6, 1),
        )
        out = StringIO()
        call_command("backfill_requisiciones_vacantes", stdout=out)

        self.assertFalse(Requisicion.objects.filter(posicion=posicion).exists())
        self.assertIn("no pasó validación", out.getvalue())

    def test_omite_una_vacante_sin_fecha_de_registro(self):
        posicion = self.create_posicion(estatus=self.estatus_vacante, tipo_requisicion=self.tipo_reemplazo)
        out = StringIO()
        call_command("backfill_requisiciones_vacantes", stdout=out)

        self.assertFalse(Requisicion.objects.filter(posicion=posicion).exists())
        self.assertIn("sin fecha de registro de vacante", out.getvalue())

    def test_es_idempotente(self):
        self.create_posicion(
            estatus=self.estatus_vacante, tipo_requisicion=self.tipo_reemplazo,
            fecha_registro_vacante=date(2025, 6, 1),
        )
        call_command("backfill_requisiciones_vacantes", stdout=StringIO())
        self.assertEqual(Requisicion.objects.count(), 1)

        out = StringIO()
        call_command("backfill_requisiciones_vacantes", stdout=out)
        self.assertEqual(Requisicion.objects.count(), 1)
        self.assertIn("Requisición(es) migrada(s): 0", out.getvalue())

    def test_no_toca_una_posicion_que_ya_tiene_una_requisicion(self):
        posicion = self.create_posicion(
            estatus=self.estatus_vacante, tipo_requisicion=self.tipo_reemplazo,
            fecha_registro_vacante=date(2025, 6, 1),
        )
        self.create_requisicion(posicion=posicion, tipo=self.tipo_reemplazo, estado=self.estado_borrador)

        call_command("backfill_requisiciones_vacantes", stdout=StringIO())

        self.assertEqual(Requisicion.objects.filter(posicion=posicion).count(), 1)
        self.assertEqual(Requisicion.objects.get(posicion=posicion).estado, self.estado_borrador)


class GenerarExcelTests(RecruitmentTestDataMixin, TestCase):
    """
    apps.recruitment.exports.generar_excel -- llena la plantilla oficial de
    GPA exacta (confirmado 2026-10-01: mismo archivo, sin rediseñarlo).
    """

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.puesto = Puesto.objects.create(name="Auxiliar de Prueba")
        cls.puesto_jefe = Puesto.objects.create(name="Supervisor de Prueba")
        unidad_level = OrganizationalLevel.objects.get(code="unidad_negocio")
        cls.unidad_node = OrganizationNode.objects.create(
            level=unidad_level, parent=cls.company_node, code="UN-TEST", name="Unidad de prueba",
        )

    def _crear_jefe_activo(self, posicion_jefe):
        persona = Persona(first_name="Ana", last_name_paternal="Jefa")
        persona.full_clean()
        persona.save()
        empleado = Empleado(persona=persona)
        empleado.full_clean()
        empleado.save()
        contrato = Contrato(empleado=empleado, posicion=posicion_jefe, fecha_ingreso=date(2020, 1, 1))
        contrato.full_clean()
        contrato.save()

    def test_genera_excel_de_reemplazo_con_los_datos_correctos(self):
        posicion_jefe = self.create_posicion(organization_node=self.unidad_node, puesto=self.puesto_jefe)
        self._crear_jefe_activo(posicion_jefe)
        posicion = self.create_posicion(
            organization_node=self.unidad_node, puesto=self.puesto, reports_to=posicion_jefe,
        )
        requisicion = self.create_requisicion(
            posicion=posicion, tipo=self.tipo_reemplazo, estado=self.estado_borrador,
            fecha_solicitud=date(2026, 1, 15), area_solicitante="Mantenimiento",
            horario_a_cubrir=HorarioACubrir.objects.get(name="7:00 - 16:00"),
            disposicion_viajar=True,
            tipo_contrato_ofrecido=TipoContratoOfrecido.objects.get(name="Planta"),
        )

        nombre_archivo, buffer = generar_excel(requisicion)

        self.assertTrue(nombre_archivo.endswith(f"_{requisicion.pk}.xlsx"))
        self.assertIn("FO-C0-CH-08", nombre_archivo)
        partes_oficiales = {
            "xl/drawings/_rels/vmlDrawing1.vml.rels",
            "xl/drawings/vmlDrawing1.vml",
            "xl/media/image1.png",
            "xl/media/image2.png",
            "xl/media/image3.png",
            "xl/media/image4.png",
            "xl/printerSettings/printerSettings1.bin",
            "xl/worksheets/_rels/sheet1.xml.rels",
        }
        with ZipFile(PLANTILLAS_DIR / "FO-C0-CH-08_reemplazo_de_personal_v4.xlsx") as original:
            contenido_original = {
                parte: original.read(parte) for parte in partes_oficiales
            }
        with ZipFile(buffer) as exportado:
            self.assertTrue(partes_oficiales.issubset(exportado.namelist()))
            for parte, contenido in contenido_original.items():
                self.assertEqual(exportado.read(parte), contenido)
            hoja_xml = exportado.read("xl/worksheets/sheet1.xml")
            self.assertIn(b'dimension ref="A1:T45"', hoja_xml)
            self.assertIn(b'row r="45"', hoja_xml)

        ws = openpyxl.load_workbook(buffer).active
        self.assertEqual(ws["D1"].value.date(), date(2026, 1, 15))
        self.assertEqual(ws["D1"].number_format, "dd/mm/yyyy")
        self.assertEqual(ws["D4"].value, self.puesto.name)
        self.assertEqual(ws["M4"].value, "Mantenimiento")
        self.assertEqual(ws["D5"].value, self.unidad_node.name)
        self.assertEqual(ws["M5"].value, self.company_node.name)
        self.assertEqual(ws["D6"].value, self.puesto_jefe.name)
        self.assertEqual(ws["M6"].value, "Jefa Ana")
        self.assertEqual(ws["I9"].value, "X")   # horario "7:00 - 16:00"
        self.assertEqual(ws["O11"].value, "X")  # disposición a viajar: sí
        self.assertIsNone(ws["Q11"].value)       # limpia el "No" premarcado de la plantilla
        self.assertEqual(ws["G17"].value, "X")  # tipo de contrato: Planta

    def test_genera_excel_de_nueva_posicion_usa_la_otra_plantilla_con_la_justificacion(self):
        posicion = self.create_posicion(puesto=self.puesto)
        requisicion = self.create_requisicion(
            posicion=posicion, tipo=self.tipo_nueva, estado=self.estado_borrador,
            justificacion="Crecimiento del área.",
        )
        nombre_archivo, buffer = generar_excel(requisicion)
        self.assertIn("FO-C0-CH-01", nombre_archivo)
        ws = openpyxl.load_workbook(buffer).active
        self.assertEqual(ws["A9"].value, "Crecimiento del área.")

    def test_campos_sin_dato_real_se_dejan_vacios_no_se_inventan(self):
        posicion = self.create_posicion()  # sin puesto, sin reports_to
        requisicion = self.create_requisicion(posicion=posicion)

        _, buffer = generar_excel(requisicion)

        ws = openpyxl.load_workbook(buffer).active
        self.assertIsNone(ws["D4"].value)  # nombre_vacante: sin Puesto capturado
        self.assertIsNone(ws["D6"].value)  # puesto_inmediato_superior: sin reports_to
        self.assertIsNone(ws["M6"].value)  # nombre_jefe_inmediato: sin reports_to
        self.assertIsNone(ws["O11"].value)  # disposición a viajar: sin dato
        self.assertIsNone(ws["Q11"].value)  # no conserva el "No" premarcado de la plantilla


class DescriptivoPuestoTestMixin(RecruitmentTestDataMixin):
    def crear_descriptivo(self, posicion, **kwargs):
        kwargs.setdefault("fecha_elaboracion", date(2026, 1, 10))
        descriptivo = DescriptivoPuesto(posicion=posicion, **kwargs)
        descriptivo.full_clean()
        descriptivo.save()
        return descriptivo

    def crear_persona(self, nombre="Luis"):
        persona = Persona(first_name=nombre, last_name_paternal="Prueba")
        persona.full_clean()
        persona.save()
        return persona


class DescriptivoVersionadoTests(DescriptivoPuestoTestMixin, TestCase):
    def test_la_primera_version_es_la_1_y_cada_copia_suma_una(self):
        posicion = self.create_posicion()
        v1 = self.crear_descriptivo(posicion)
        self.assertEqual(v1.version, 1)
        v1.congelar()
        v2 = copiar_version(v1)
        self.assertEqual(v2.version, 2)
        v2.congelar()
        self.assertEqual(copiar_version(v2).version, 3)

    def test_las_versiones_se_numeran_por_posicion(self):
        otra = self.create_posicion()
        self.assertEqual(self.crear_descriptivo(self.create_posicion()).version, 1)
        self.assertEqual(self.crear_descriptivo(otra).version, 1)

    def test_no_se_puede_abrir_un_segundo_borrador_en_la_misma_posicion(self):
        posicion = self.create_posicion()
        self.crear_descriptivo(posicion)
        segundo = DescriptivoPuesto(posicion=posicion, fecha_elaboracion=date(2026, 2, 1))
        with self.assertRaises(ValidationError) as contexto:
            segundo.full_clean()
        self.assertIn("posicion", contexto.exception.message_dict)

    def test_la_base_tambien_impide_dos_borradores_aunque_se_salten_las_validaciones(self):
        posicion = self.create_posicion()
        self.crear_descriptivo(posicion)
        with self.assertRaises(IntegrityError), transaction.atomic():
            DescriptivoPuesto.objects.create(posicion=posicion, fecha_elaboracion=date(2026, 2, 1))

    def test_congelar_sella_la_version_y_ya_no_se_puede_modificar(self):
        descriptivo = self.crear_descriptivo(self.create_posicion(), proposito="Original")
        descriptivo.congelar()
        self.assertTrue(descriptivo.esta_congelado)

        # Ni con el mismo objeto, ni con uno recién leído de la base.
        for objeto in (descriptivo, DescriptivoPuesto.objects.get(pk=descriptivo.pk)):
            objeto.proposito = "Cambiado"
            with self.assertRaises(ValidationError):
                objeto.save()
            with self.assertRaises(ValidationError):
                objeto.full_clean()
        self.assertEqual(DescriptivoPuesto.objects.get(pk=descriptivo.pk).proposito, "Original")

    def test_un_objeto_viejo_en_memoria_no_pasa_por_encima_de_un_congelado_reciente(self):
        descriptivo = self.crear_descriptivo(self.create_posicion(), proposito="Original")
        copia_vieja = DescriptivoPuesto.objects.get(pk=descriptivo.pk)  # aún sin congelar en memoria
        descriptivo.congelar()
        copia_vieja.proposito = "Cambio tardío"
        with self.assertRaises(ValidationError):
            copia_vieja.save()

    def test_no_se_congela_dos_veces(self):
        descriptivo = self.crear_descriptivo(self.create_posicion())
        descriptivo.congelar()
        with self.assertRaises(ValidationError):
            descriptivo.congelar()

    def test_una_version_congelada_no_se_borra_pero_un_borrador_si(self):
        posicion = self.create_posicion()
        borrador = self.crear_descriptivo(posicion)
        borrador.delete()
        self.assertFalse(DescriptivoPuesto.objects.filter(pk=borrador.pk).exists())
        self.assertTrue(DescriptivoPuesto.all_objects.filter(pk=borrador.pk, is_deleted=True).exists())

        congelado = self.crear_descriptivo(posicion)
        congelado.congelar()
        with self.assertRaises(ValidationError):
            congelado.delete()
        self.assertTrue(DescriptivoPuesto.objects.filter(pk=congelado.pk).exists())

    def test_vigente_de_es_la_congelada_mas_reciente_e_ignora_borradores(self):
        posicion = self.create_posicion()
        self.assertIsNone(DescriptivoPuesto.vigente_de(posicion))
        v1 = self.crear_descriptivo(posicion)
        self.assertIsNone(DescriptivoPuesto.vigente_de(posicion))  # un borrador no es vigente
        v1.congelar()
        self.assertEqual(DescriptivoPuesto.vigente_de(posicion), v1)
        v2 = copiar_version(v1)  # borrador: la vigente sigue siendo v1
        self.assertEqual(DescriptivoPuesto.vigente_de(posicion), v1)
        v2.congelar()
        self.assertEqual(DescriptivoPuesto.vigente_de(posicion), v2)

    def test_congelado_bloquea_funciones_indicadores_y_casillas(self):
        descriptivo = self.crear_descriptivo(self.create_posicion())
        funcion = FuncionPuesto.objects.create(descriptivo=descriptivo, orden=1, texto="Supervisar")
        indicador = IndicadorDesempeno.objects.create(descriptivo=descriptivo, orden=1, texto="Meta diaria")
        descriptivo.competencias.add(CompetenciaConductual.objects.get(name="Liderazgo"))
        descriptivo.congelar()

        with self.assertRaises(ValidationError):
            FuncionPuesto.objects.create(descriptivo=descriptivo, orden=2, texto="Otra")
        funcion.texto = "Cambiada"
        with self.assertRaises(ValidationError):
            funcion.save()
        with self.assertRaises(ValidationError):
            funcion.delete()
        with self.assertRaises(ValidationError):
            IndicadorDesempeno.objects.create(descriptivo=descriptivo, orden=2, texto="Otro")
        with self.assertRaises(ValidationError):
            indicador.delete()
        # Las casillas (M2M) se tocan dentro de un atomic() propio de Django,
        # sin savepoint: sin este atomic() de la prueba, la excepción dejaría
        # rota la transacción del TestCase para todo lo que sigue.
        with self.assertRaises(ValidationError), transaction.atomic():
            descriptivo.competencias.add(CompetenciaConductual.objects.get(name="Innovación"))
        with self.assertRaises(ValidationError), transaction.atomic():
            descriptivo.competencias.clear()
        with self.assertRaises(ValidationError), transaction.atomic():
            descriptivo.recursos.set([RecursoAsignado.objects.get(name="Uniforme/EPP")])

        self.assertEqual(descriptivo.funciones.count(), 1)
        self.assertEqual(descriptivo.indicadores.count(), 1)
        self.assertEqual(list(descriptivo.competencias.values_list("name", flat=True)), ["Liderazgo"])
        self.assertEqual(descriptivo.recursos.count(), 0)

    def test_el_borrador_si_se_edita_libremente(self):
        descriptivo = self.crear_descriptivo(self.create_posicion())
        FuncionPuesto.objects.create(descriptivo=descriptivo, orden=1, texto="Supervisar")
        descriptivo.recursos.add(RecursoAsignado.objects.get(name="Uniforme/EPP"))
        descriptivo.proposito = "Nuevo propósito"
        descriptivo.full_clean()
        descriptivo.save()
        self.assertEqual(descriptivo.funciones.count(), 1)
        self.assertEqual(descriptivo.recursos.count(), 1)

    def test_el_orden_de_funciones_no_se_repite_dentro_de_un_descriptivo(self):
        descriptivo = self.crear_descriptivo(self.create_posicion())
        FuncionPuesto.objects.create(descriptivo=descriptivo, orden=1, texto="Una")
        with self.assertRaises(IntegrityError), transaction.atomic():
            FuncionPuesto.objects.create(descriptivo=descriptivo, orden=1, texto="Repetida")


def _abrir_word(buffer):
    with ZipFile(buffer) as paquete:
        return minidom.parseString(paquete.read("word/document.xml"))


def _plantilla_word():
    with ZipFile(PLANTILLA_WORD) as paquete:
        return minidom.parseString(paquete.read("word/document.xml"))


def _casillas_marcadas(documento):
    marcadas = set()
    for indice, sdt in enumerate(documento.getElementsByTagName("w:sdt")):
        for casilla in sdt.getElementsByTagName("w14:checked"):
            if casilla.getAttribute("w14:val") == "1":
                marcadas.add(indice)
    return marcadas


def _controles_cambiados(documento):
    """Posicion de los controles que ya no son iguales a los de la plantilla oficial."""
    base = list(_plantilla_word().getElementsByTagName("w:sdt"))
    nuevos = list(documento.getElementsByTagName("w:sdt"))
    assert len(base) == len(nuevos), "se agregaron o quitaron controles"
    return {i for i, (a, b) in enumerate(zip(base, nuevos)) if a.toxml() != b.toxml()}


class GenerarWordTests(DescriptivoPuestoTestMixin, TestCase):
    """
    apps.recruitment.exports_word.generar_word -- llena el Descriptivo oficial
    de GPA (FO-C0-CH-04) sin rediseñarlo. La prueba central: solo cambian los
    controles que tienen dato; todo lo demás queda igual a la plantilla.
    """

    def _descriptivo(self, **kwargs):
        return self.crear_descriptivo(self.create_posicion(), **kwargs)

    def test_conserva_intacto_todo_el_paquete_salvo_el_documento(self):
        nombre_archivo, buffer = generar_word(self._descriptivo(nombre_puesto="Soldador"))

        self.assertIn("FO-C0-CH-04", nombre_archivo)
        self.assertTrue(nombre_archivo.endswith(".docx"))
        with ZipFile(PLANTILLA_WORD) as original, ZipFile(buffer) as exportado:
            self.assertEqual(original.namelist(), exportado.namelist())
            for parte in original.namelist():
                if parte == "word/document.xml":
                    continue
                self.assertEqual(exportado.read(parte), original.read(parte), parte)
            self.assertIn("word/media/image1.png", exportado.namelist())  # imagenes del encabezado
            self.assertIn("word/glossary/document.xml", exportado.namelist())  # textos guia

    def test_un_descriptivo_vacio_solo_cambia_la_fecha_y_no_marca_nada(self):
        _, buffer = generar_word(self._descriptivo())
        documento = _abrir_word(buffer)
        self.assertEqual(_controles_cambiados(documento), {5})
        self.assertEqual(_casillas_marcadas(documento), set())

    def test_llena_textos_y_fecha_y_deja_igual_todo_lo_demas(self):
        _, buffer = generar_word(self._descriptivo(
            nombre_puesto="Soldador", empresa="GPA Advanced Manufacturing", area_departamento="Mass Production",
            reporta_a="Supervisor", proposito="Soldar estructuras", decisiones_operativas="Asignar turnos",
            relaciones_externas="Proveedores", escolaridad_minima="Secundaria",
            fecha_elaboracion=date(2026, 1, 10),
        ))
        documento = _abrir_word(buffer)
        controles = list(documento.getElementsByTagName("w:sdt"))

        # Solo cambian los controles con dato (supervisa_a quedo vacio -> intacto).
        self.assertEqual(_controles_cambiados(documento), {0, 1, 2, 3, 5, 21, 28, 32, 33})
        self.assertEqual(_texto_de(controles[0]), "Soldador")
        self.assertEqual(_texto_de(controles[21]), "Soldar estructuras")
        self.assertEqual(_texto_de(controles[5]), "10 de enero de 2026")
        self.assertEqual(
            controles[5].getElementsByTagName("w:date")[0].getAttribute("w:fullDate"), "2026-01-10T00:00:00Z",
        )
        for indice in (0, 1, 2, 3, 5, 21, 28, 32, 33):
            self.assertFalse(controles[indice].getElementsByTagName("w:showingPlcHdr"), indice)
        for indice in (4, 27, 29, 49):  # sin dato, o texto del formato: se quedan con su texto guia
            self.assertEqual(
                controles[indice].toxml(),
                list(_plantilla_word().getElementsByTagName("w:sdt"))[indice].toxml(),
            )
        # Escrito con el estilo "Cuerpo" de la plantilla, no con la cursiva gris del texto guia.
        run = controles[0].getElementsByTagName("w:r")[0]
        self.assertEqual(run.getElementsByTagName("w:rStyle")[0].getAttribute("w:val"), "Cuerpo")
        self.assertFalse(run.getElementsByTagName("w:i"))
        self.assertFalse(run.getElementsByTagName("w:color"))

    def test_marca_exactamente_las_casillas_que_corresponden(self):
        descriptivo = self._descriptivo(
            edad=RangoEdad.objects.get(name="26-35 años"), disponibilidad_viajar=True,
            dias_por_laborar=DiasPorLaborar.objects.get(name="Lunes a Viernes"),
            horario=HorarioACubrir.objects.get(name="7:00 - 16:00"),
        )
        descriptivo.competencias.add(*CompetenciaConductual.objects.filter(name__in=["Liderazgo", "Integridad"]))
        descriptivo.recursos.add(*RecursoAsignado.objects.filter(
            name__in=["Uniforme/EPP", "Vehículo asignado", "Fondo fijo o caja chica asignada (para manejo de efectivo)"],
        ))

        _, buffer = generar_word(descriptivo)
        documento = _abrir_word(buffer)

        # 7 edad 26-35 | 11 viajar SI | 13 Lunes a Viernes | 17 07:00-16:00
        # 43 Liderazgo, 37 Integridad | 59 Uniforme/EPP, 51 Vehiculo, 57 Fondo fijo
        self.assertEqual(_casillas_marcadas(documento), {7, 11, 13, 17, 43, 37, 59, 51, 57})
        self.assertEqual(_controles_cambiados(documento), {5, 7, 11, 13, 17, 43, 37, 59, 51, 57})
        marcada = list(documento.getElementsByTagName("w:sdt"))[7]
        self.assertEqual(_texto_de(marcada), "☒")
        simbolo = marcada.getElementsByTagName("w:sdtContent")[0]  # no las propiedades del control (esas son Arial)
        fuentes = simbolo.getElementsByTagName("w:rFonts")[0]
        self.assertEqual(fuentes.getAttribute("w:ascii"), "MS Gothic")  # la fuente del estado "marcado"
        self.assertEqual(fuentes.getAttribute("w:eastAsia"), "MS Gothic")
        self.assertFalse(fuentes.hasAttribute("w:cs"))

    def test_viajar_no_marca_si_no_hay_dato_y_marca_no_cuando_es_falso(self):
        _, sin_dato = generar_word(self._descriptivo())
        self.assertEqual(_casillas_marcadas(_abrir_word(sin_dato)), set())
        _, con_no = generar_word(self._descriptivo(disponibilidad_viajar=False))
        self.assertEqual(_casillas_marcadas(_abrir_word(con_no)), {12})

    def test_renglones_otro_se_llenan_subrayados_y_los_demas_conservan_su_linea(self):
        descriptivo = self._descriptivo(
            edad=RangoEdad.objects.get(name="Otro"), edad_otro="52 años",
            horario=HorarioACubrir.objects.get(name="Otro"), horario_otro="Turno rolado",
            competencias_otras="Puntualidad", recursos_otro="Radio de comunicación",
        )
        _, buffer = generar_word(descriptivo)
        documento = _abrir_word(buffer)
        controles = list(documento.getElementsByTagName("w:sdt"))

        edad = _texto_de(_ancestro(controles[10], "w:p"))
        self.assertIn("Otro: 52 años", edad)
        self.assertNotIn("_", edad)
        self.assertTrue(_ancestro(controles[10], "w:p").getElementsByTagName("w:u"))  # el valor va subrayado
        self.assertIn("Otro: Turno rolado", _texto_de(_ancestro(controles[20], "w:p")))
        self.assertIn("____", _texto_de(_ancestro(controles[15], "w:p")))  # dias "Otro": sin dato, conserva su linea
        otras = [p for p in documento.getElementsByTagName("w:p") if _texto_de(p).startswith("Otras:")]
        self.assertEqual([_texto_de(p) for p in otras], ["Otras: Puntualidad"])
        fila_uniforme = _ancestro(controles[59], "w:tr")
        self.assertIn("Otro: Radio de comunicación", _texto_de(fila_uniforme))
        self.assertNotIn("_", _texto_de(fila_uniforme))
        self.assertEqual(_casillas_marcadas(documento), {10, 20})

    def test_funciones_e_indicadores_van_en_su_fila_numerada(self):
        descriptivo = self._descriptivo()
        for numero in (1, 2):
            FuncionPuesto.objects.create(descriptivo=descriptivo, orden=numero, texto=f"Función {numero}")
        IndicadorDesempeno.objects.create(descriptivo=descriptivo, orden=1, texto="Piezas por turno")

        _, buffer = generar_word(descriptivo)
        documento = _abrir_word(buffer)
        controles = list(documento.getElementsByTagName("w:sdt"))

        self.assertEqual(_controles_cambiados(documento), {5, 22, 23, 60})
        self.assertEqual(_texto_de(controles[22]), "Función 1")
        self.assertEqual(_texto_de(controles[23]), "Función 2")
        self.assertEqual(_texto_de(controles[60]), "Piezas por turno")
        # Las filas sin dato (3 a 5) conservan el texto guia de la plantilla.
        self.assertTrue(controles[24].getElementsByTagName("w:showingPlcHdr"))

    def test_mas_filas_de_las_que_trae_el_formulario_clona_la_ultima(self):
        descriptivo = self._descriptivo()
        for numero in range(1, 8):
            FuncionPuesto.objects.create(descriptivo=descriptivo, orden=numero, texto=f"Función {numero}")
        for numero in range(1, 5):
            IndicadorDesempeno.objects.create(descriptivo=descriptivo, orden=numero, texto=f"Indicador texto {numero}")

        _, buffer = generar_word(descriptivo)
        documento = _abrir_word(buffer)

        self.assertEqual(len(documento.getElementsByTagName("w:sdt")), 63 + 2 + 1)
        filas = [tr for tr in documento.getElementsByTagName("w:tr")]
        responsabilidades = [_texto_de(tr) for tr in filas if _texto_de(tr).startswith("Responsabilidad ")]
        self.assertEqual(
            responsabilidades, [f"Responsabilidad {n}Función {n}" for n in range(1, 8)],
        )
        indicadores = [_texto_de(tr) for tr in filas if _texto_de(tr).startswith("Indicador ")]
        self.assertEqual(indicadores, [f"Indicador {n}Indicador texto {n}" for n in range(1, 5)])
        # Cada control nuevo trae su propio id y los parrafos clonados no repiten paraId.
        ids = [e.getAttribute("w:val") for e in documento.getElementsByTagName("w:id") if e.parentNode.tagName == "w:sdtPr"]
        self.assertEqual(len(ids), len(set(ids)))
        para_ids = [e.getAttribute("w14:paraId") for e in documento.getElementsByTagName("*") if e.hasAttribute("w14:paraId")]
        self.assertEqual(len(para_ids), len(set(para_ids)))
        # Las filas clonadas quedan justo despues de la ultima original, no al final del documento.
        self.assertEqual(_texto_de(filas[[i for i, tr in enumerate(filas) if _texto_de(tr).startswith("Responsabilidad 5")][0] + 1])[:16], "Responsabilidad ")

    def test_caracteres_especiales_y_saltos_de_linea_no_rompen_el_documento(self):
        _, buffer = generar_word(self._descriptivo(proposito="Tom & Jerry <b>\"ñandú\"</b>\nSegunda línea\r\nTercera\x07"))
        documento = _abrir_word(buffer)  # si no fuera XML valido, esto falla
        control = list(documento.getElementsByTagName("w:sdt"))[21]
        self.assertEqual(_texto_de(control), "Tom & Jerry <b>\"ñandú\"</b>Segunda líneaTercera")
        # Un parrafo por linea (con el mismo formato) y SIN saltos manuales: en una
        # celda justificada, Word estira hasta el margen la linea que termina en w:br.
        parrafos = control.getElementsByTagName("w:p")
        self.assertEqual([_texto_de(p) for p in parrafos], ["Tom & Jerry <b>\"ñandú\"</b>", "Segunda línea", "Tercera"])
        self.assertFalse(control.getElementsByTagName("w:br"))
        self.assertEqual(len({p.getElementsByTagName("w:jc")[0].getAttribute("w:val") for p in parrafos}), 1)

    def test_el_valor_de_un_renglon_otro_es_de_una_sola_linea(self):
        _, buffer = generar_word(self._descriptivo(
            edad=RangoEdad.objects.get(name="Otro"), edad_otro="52 años\ny medio",
        ))
        controles = list(_abrir_word(buffer).getElementsByTagName("w:sdt"))
        self.assertIn("Otro: 52 años y medio", _texto_de(_ancestro(controles[10], "w:p")))

    def test_las_firmas_no_se_tocan_aunque_haya_conformidades(self):
        descriptivo = self._descriptivo(nombre_puesto="Soldador")
        descriptivo.congelar()
        ConformidadDescriptivo.objects.create(
            descriptivo=descriptivo, rol=RolConformidad.objects.get(name="Jefe inmediato"),
            fecha=date(2026, 3, 1), nombre_manual="Ana Jefa",
        )
        _, buffer = generar_word(descriptivo)
        filas_nuevas = list(_abrir_word(buffer).getElementsByTagName("w:tr"))[-3:]
        filas_plantilla = list(_plantilla_word().getElementsByTagName("w:tr"))[-3:]
        self.assertEqual([f.toxml() for f in filas_nuevas], [f.toxml() for f in filas_plantilla])
        self.assertNotIn("Ana Jefa", "".join(_texto_de(f) for f in filas_nuevas))

    def test_los_codigos_sembrados_coinciden_con_los_del_mapa(self):
        for modelo, mapa in (
            (RangoEdad, exports_word._EDAD), (DiasPorLaborar, exports_word._DIAS),
            (HorarioACubrir, exports_word._HORARIO), (CompetenciaConductual, exports_word._COMPETENCIAS),
            (RecursoAsignado, exports_word._RECURSOS),
        ):
            with self.subTest(modelo=modelo.__name__):
                self.assertEqual(set(modelo.objects.values_list("code", flat=True)), set(mapa))

    def test_si_la_plantilla_cambia_falla_con_un_mensaje_claro_en_vez_de_escribir_en_otro_campo(self):
        descriptivo = self._descriptivo(empresa="GPA")
        with mock.patch.dict(exports_word._TEXTOS, {"empresa": (1, "Una etiqueta que ya no existe")}):
            with self.assertRaises(RuntimeError) as contexto:
                generar_word(descriptivo)
        self.assertIn("La plantilla del Descriptivo cambió", str(contexto.exception))
        with mock.patch.object(exports_word, "_TOTAL_CONTROLES", 62):
            with self.assertRaises(RuntimeError):
                generar_word(descriptivo)


class DescriptivoServiciosTests(DescriptivoPuestoTestMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.puesto = Puesto.objects.create(name="Auxiliar de Prueba")
        cls.puesto_jefe = Puesto.objects.create(name="Supervisor de Prueba")

    def test_crear_borrador_precarga_lo_que_el_sistema_ya_sabe(self):
        jefe = self.create_posicion(puesto=self.puesto_jefe)
        posicion = self.create_posicion(puesto=self.puesto, reports_to=jefe)

        descriptivo = crear_borrador(posicion)

        self.assertEqual(descriptivo.version, 1)
        self.assertFalse(descriptivo.esta_congelado)
        self.assertEqual(descriptivo.nombre_puesto, "Auxiliar de Prueba")
        self.assertEqual(descriptivo.empresa, self.company_node.name)
        self.assertEqual(descriptivo.area_departamento, self.company_node.name)
        self.assertEqual(descriptivo.reporta_a, "Supervisor de Prueba")
        self.assertEqual(descriptivo.supervisa_a, "")  # no se precarga: cobertura incompleta de reports_to

    def test_crear_borrador_deja_vacio_lo_que_no_se_sabe_en_vez_de_inventarlo(self):
        descriptivo = crear_borrador(self.create_posicion())  # sin Puesto ni reports_to
        self.assertEqual(descriptivo.nombre_puesto, "")
        self.assertEqual(descriptivo.reporta_a, "")

    def test_crear_borrador_respeta_la_regla_de_un_solo_borrador(self):
        posicion = self.create_posicion()
        crear_borrador(posicion)
        with self.assertRaises(ValidationError):
            crear_borrador(posicion)

    def test_crear_borrador_registra_quien_lo_creo(self):
        usuario = get_user_model().objects.create_user(
            username="redactor", email="redactor@example.com", password="strong-test-password",
        )
        descriptivo = crear_borrador(self.create_posicion(), user=usuario)
        self.assertEqual(descriptivo.created_by, usuario)

    def test_copiar_version_trae_todo_el_contenido_sin_tocar_la_original(self):
        posicion = self.create_posicion()
        original = self.crear_descriptivo(
            posicion, nombre_puesto="Soldador", proposito="Soldar", escolaridad_minima="Secundaria",
            edad=RangoEdad.objects.get(name="26-35 años"), disponibilidad_viajar=True,
            dias_por_laborar=DiasPorLaborar.objects.get(name="Lunes a Viernes"),
            horario=HorarioACubrir.objects.get(name="7:00 - 16:00"), competencias_otras="Puntualidad",
        )
        FuncionPuesto.objects.create(descriptivo=original, orden=1, texto="Soldar piezas")
        FuncionPuesto.objects.create(descriptivo=original, orden=2, texto="Revisar acabados")
        IndicadorDesempeno.objects.create(descriptivo=original, orden=1, texto="Piezas por turno")
        original.competencias.add(*CompetenciaConductual.objects.filter(name__in=["Liderazgo", "Integridad"]))
        original.recursos.add(RecursoAsignado.objects.get(name="Uniforme/EPP"))
        original.congelar()

        copia = copiar_version(original)

        self.assertNotEqual(copia.pk, original.pk)
        self.assertEqual(copia.version, 2)
        self.assertFalse(copia.esta_congelado)  # el borrador nuevo sí se puede editar
        self.assertEqual(copia.nombre_puesto, "Soldador")
        self.assertEqual(copia.proposito, "Soldar")
        self.assertEqual(copia.edad.name, "26-35 años")
        self.assertIs(copia.disponibilidad_viajar, True)
        self.assertEqual(copia.horario.name, "7:00 - 16:00")
        self.assertEqual(copia.competencias_otras, "Puntualidad")
        self.assertEqual(
            list(copia.funciones.values_list("orden", "texto")),
            [(1, "Soldar piezas"), (2, "Revisar acabados")],
        )
        self.assertEqual(list(copia.indicadores.values_list("texto", flat=True)), ["Piezas por turno"])
        self.assertEqual(set(copia.competencias.values_list("name", flat=True)), {"Liderazgo", "Integridad"})
        self.assertEqual(list(copia.recursos.values_list("name", flat=True)), ["Uniforme/EPP"])

        # Editar la copia no toca la original congelada.
        copia.proposito = "Soldar y supervisar"
        copia.full_clean()
        copia.save()
        FuncionPuesto.objects.create(descriptivo=copia, orden=3, texto="Capacitar")
        self.assertEqual(DescriptivoPuesto.objects.get(pk=original.pk).proposito, "Soldar")
        self.assertEqual(original.funciones.count(), 2)

    def test_copiar_version_no_copia_las_conformidades(self):
        posicion = self.create_posicion()
        original = self.crear_descriptivo(posicion)
        original.congelar()
        ConformidadDescriptivo.objects.create(
            descriptivo=original, rol=RolConformidad.objects.get(name="Jefe inmediato"),
            fecha=date(2026, 2, 1), nombre_manual="Ana Jefa",
        )
        copia = copiar_version(original)
        self.assertEqual(copia.conformidades.count(), 0)

    def test_copiar_version_falla_si_ya_hay_un_borrador_abierto(self):
        posicion = self.create_posicion()
        v1 = self.crear_descriptivo(posicion)
        v1.congelar()
        copiar_version(v1)
        with self.assertRaises(ValidationError):
            copiar_version(v1)


class ConformidadDescriptivoTests(DescriptivoPuestoTestMixin, TestCase):
    def _congelado(self):
        descriptivo = self.crear_descriptivo(self.create_posicion())
        descriptivo.congelar()
        return descriptivo

    def _conformidad(self, descriptivo, rol, **kwargs):
        conformidad = ConformidadDescriptivo(
            descriptivo=descriptivo, rol=RolConformidad.objects.get(name=rol), **kwargs,
        )
        conformidad.full_clean()
        conformidad.save()
        return conformidad

    def test_no_se_da_conformidad_sobre_un_borrador(self):
        borrador = self.crear_descriptivo(self.create_posicion())
        conformidad = ConformidadDescriptivo(
            descriptivo=borrador, rol=RolConformidad.objects.get(name="Jefe inmediato"),
        )
        with self.assertRaises(ValidationError) as contexto:
            conformidad.full_clean()
        self.assertIn("descriptivo", contexto.exception.message_dict)

    def test_las_conformidades_se_registran_despues_de_congelar_sin_romper_el_congelado(self):
        descriptivo = self._congelado()
        self._conformidad(descriptivo, "Jefe inmediato", fecha=date(2026, 3, 1), nombre_manual="Ana Jefa")
        self._conformidad(descriptivo, "Capital Humano", fecha=date(2026, 3, 2), nombre_manual="Rosa RH")
        self.assertEqual(descriptivo.conformidades.count(), 2)
        self.assertTrue(DescriptivoPuesto.objects.get(pk=descriptivo.pk).esta_congelado)

    def test_el_colaborador_exige_una_persona_especifica(self):
        descriptivo = self._congelado()
        conformidad = ConformidadDescriptivo(
            descriptivo=descriptivo, rol=RolConformidad.objects.get(name="Colaborador"),
            fecha=date(2026, 3, 1), nombre_manual="Alguien",
        )
        with self.assertRaises(ValidationError) as contexto:
            conformidad.full_clean()
        self.assertIn("persona", contexto.exception.message_dict)

    def test_dos_personas_distintas_pueden_dar_su_conformidad_como_colaborador(self):
        descriptivo = self._congelado()
        self._conformidad(descriptivo, "Colaborador", persona=self.crear_persona("Luis"), fecha=date(2026, 3, 1))
        self._conformidad(descriptivo, "Colaborador", persona=self.crear_persona("Mara"), fecha=date(2027, 3, 1))
        self.assertEqual(descriptivo.conformidades.count(), 2)

    def test_la_misma_persona_no_firma_dos_veces_la_misma_version(self):
        descriptivo = self._congelado()
        persona = self.crear_persona()
        self._conformidad(descriptivo, "Colaborador", persona=persona, fecha=date(2026, 3, 1))
        repetida = ConformidadDescriptivo(
            descriptivo=descriptivo, rol=RolConformidad.objects.get(name="Colaborador"), persona=persona,
        )
        with self.assertRaises(ValidationError):
            repetida.full_clean()

    def test_un_rol_sin_persona_se_registra_una_sola_vez_por_version(self):
        descriptivo = self._congelado()
        self._conformidad(descriptivo, "Jefe inmediato", fecha=date(2026, 3, 1), nombre_manual="Ana")
        repetida = ConformidadDescriptivo(
            descriptivo=descriptivo, rol=RolConformidad.objects.get(name="Jefe inmediato"),
        )
        with self.assertRaises(ValidationError):
            repetida.full_clean()

    def test_con_fecha_debe_decir_quien_dio_la_conformidad(self):
        descriptivo = self._congelado()
        conformidad = ConformidadDescriptivo(
            descriptivo=descriptivo, rol=RolConformidad.objects.get(name="Capital Humano"), fecha=date(2026, 3, 1),
        )
        with self.assertRaises(ValidationError) as contexto:
            conformidad.full_clean()
        self.assertIn("nombre_manual", contexto.exception.message_dict)

    def test_sin_fecha_queda_pendiente_sin_exigir_quien(self):
        descriptivo = self._congelado()
        pendiente = self._conformidad(descriptivo, "Capital Humano")
        self.assertIsNone(pendiente.fecha)

    def test_conformidad_aceptada_dentro_del_sistema_usa_usuario(self):
        descriptivo = self._congelado()
        usuario = get_user_model().objects.create_user(
            username="jefe-sistema", email="jefe-sistema@example.com", password="strong-test-password",
        )
        conformidad = self._conformidad(descriptivo, "Jefe inmediato", fecha=date(2026, 3, 1), usuario=usuario)
        self.assertEqual(conformidad.usuario, usuario)


class DescriptivoPuestoAPITests(DescriptivoPuestoTestMixin, APITestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        user_model = get_user_model()
        cls.gestor = user_model.objects.create_user(
            username="descriptivo-gestor", email="descriptivo-gestor@example.com",
            password="strong-test-password", role=UserRole.objects.get(code="capital-humano"),
        )
        cls.colaborador = user_model.objects.create_user(
            username="descriptivo-colaborador", email="descriptivo-colaborador@example.com",
            password="strong-test-password", role=UserRole.objects.get(code="colaborador"),
        )
        cls.puesto = Puesto.objects.create(name="Soldador de Prueba")

    def _payload(self, posicion, **extra):
        return {"posicion": str(posicion.pk), "fecha_elaboracion": "2026-01-10", **extra}

    def _crear_por_api(self, posicion, **extra):
        self.client.force_authenticate(user=self.gestor)
        response = self.client.post(reverse("descriptivopuesto-list"), self._payload(posicion, **extra), format="json")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        return response.data

    def test_catalogos_son_de_solo_lectura_para_cualquier_autenticado(self):
        self.client.force_authenticate(user=self.colaborador)
        for nombre, cuantos in (
            ("rangoedad-list", 5), ("diasporlaborar-list", 3), ("competenciaconductual-list", 12),
            ("recursoasignado-list", 10), ("rolconformidad-list", 3),
        ):
            with self.subTest(catalogo=nombre):
                self.assertEqual(self.client.get(reverse(nombre)).data["count"], cuantos)
                self.assertEqual(
                    self.client.post(reverse(nombre), {"name": "Nuevo"}, format="json").status_code,
                    status.HTTP_405_METHOD_NOT_ALLOWED,
                )

    def test_la_lista_trae_la_etiqueta_de_la_posicion_y_busca_y_ordena(self):
        puesto_a = Puesto.objects.create(name="Auxiliar de Almacén")
        posicion_a = self.create_posicion(puesto=puesto_a)
        posicion_b = self.create_posicion(puesto=self.puesto)
        uno = self.crear_descriptivo(posicion_a, nombre_puesto="Auxiliar de Almacén", fecha_elaboracion=date(2026, 2, 1))
        dos = self.crear_descriptivo(posicion_b, nombre_puesto="Soldador de Prueba", fecha_elaboracion=date(2026, 4, 1))
        self.client.force_authenticate(user=self.colaborador)
        url = reverse("descriptivopuesto-list")

        datos = {r["id"]: r for r in self.client.get(url).data["results"]}
        self.assertEqual(datos[str(uno.pk)]["posicion_etiqueta"], "Auxiliar de Almacén — Empresa de prueba")

        # Sin acentos: "almacen" encuentra "Almacén".
        encontrados = self.client.get(url, {"search": "almacen"}).data["results"]
        self.assertEqual([r["id"] for r in encontrados], [str(uno.pk)])

        ordenados = self.client.get(url, {"ordering": "-fecha_elaboracion"}).data["results"]
        self.assertEqual([r["id"] for r in ordenados], [str(dos.pk), str(uno.pk)])

    def test_un_colaborador_puede_leer_pero_no_escribir_descriptivos(self):
        posicion = self.create_posicion()
        descriptivo = self.crear_descriptivo(posicion)
        self.client.force_authenticate(user=self.colaborador)

        self.assertEqual(self.client.get(reverse("descriptivopuesto-list")).data["count"], 1)
        self.assertEqual(
            self.client.get(reverse("descriptivopuesto-detail", args=[descriptivo.pk])).status_code, status.HTTP_200_OK,
        )
        self.assertEqual(
            self.client.post(reverse("descriptivopuesto-list"), self._payload(self.create_posicion()), format="json").status_code,
            status.HTTP_403_FORBIDDEN,
        )
        self.assertEqual(
            self.client.patch(
                reverse("descriptivopuesto-detail", args=[descriptivo.pk]), {"proposito": "x"}, format="json",
            ).status_code,
            status.HTTP_403_FORBIDDEN,
        )
        self.assertEqual(
            self.client.delete(reverse("descriptivopuesto-detail", args=[descriptivo.pk])).status_code,
            status.HTTP_403_FORBIDDEN,
        )
        for accion, metodo in (("crear-borrador", None), ("copiar", descriptivo.pk), ("congelar", descriptivo.pk)):
            with self.subTest(accion=accion):
                url = reverse(f"descriptivopuesto-{accion}", args=[metodo] if metodo else [])
                self.assertEqual(self.client.post(url, {}, format="json").status_code, status.HTTP_403_FORBIDDEN)

    def test_gestor_crea_un_descriptivo_completo_en_una_sola_llamada(self):
        posicion = self.create_posicion(puesto=self.puesto)
        liderazgo = CompetenciaConductual.objects.get(name="Liderazgo")
        epp = RecursoAsignado.objects.get(name="Uniforme/EPP")
        datos = self._crear_por_api(
            posicion,
            nombre_puesto="Soldador", proposito="Soldar estructuras",
            edad=RangoEdad.objects.get(name="26-35 años").pk, disponibilidad_viajar=False,
            horario=HorarioACubrir.objects.get(name="7:00 - 16:00").pk,
            funciones=[{"texto": "Soldar piezas"}, {"texto": "Revisar acabados"}],
            indicadores=[{"texto": "Piezas por turno"}],
            competencias=[liderazgo.pk], recursos=[epp.pk],
        )

        self.assertEqual(datos["version"], 1)
        self.assertFalse(datos["esta_congelado"])
        self.assertEqual(datos["funciones"], [{"orden": 1, "texto": "Soldar piezas"}, {"orden": 2, "texto": "Revisar acabados"}])
        self.assertEqual(datos["indicadores"], [{"orden": 1, "texto": "Piezas por turno"}])
        self.assertEqual(datos["competencias"], [liderazgo.pk])
        self.assertEqual(datos["recursos"], [epp.pk])
        descriptivo = DescriptivoPuesto.objects.get(pk=datos["id"])
        self.assertEqual(descriptivo.created_by, self.gestor)
        self.assertEqual(descriptivo.funciones.first().created_by, self.gestor)
        self.assertIs(descriptivo.disponibilidad_viajar, False)

    def test_crear_con_listas_invalidas_no_deja_nada_a_medias(self):
        posicion = self.create_posicion()
        self.client.force_authenticate(user=self.gestor)
        response = self.client.post(
            reverse("descriptivopuesto-list"),
            self._payload(posicion, funciones=[{"texto": "Valida"}, {"texto": ""}]), format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(DescriptivoPuesto.objects.count(), 0)
        self.assertEqual(FuncionPuesto.objects.count(), 0)

    def test_patch_reemplaza_las_listas_que_se_mandan_y_respeta_las_que_se_omiten(self):
        posicion = self.create_posicion()
        liderazgo = CompetenciaConductual.objects.get(name="Liderazgo")
        integridad = CompetenciaConductual.objects.get(name="Integridad")
        datos = self._crear_por_api(
            posicion, funciones=[{"texto": "Una"}, {"texto": "Dos"}], indicadores=[{"texto": "Meta"}],
            competencias=[liderazgo.pk],
        )
        url = reverse("descriptivopuesto-detail", args=[datos["id"]])

        response = self.client.patch(
            url, {"funciones": [{"texto": "Nueva uno"}], "competencias": [integridad.pk]}, format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        self.assertEqual(response.data["funciones"], [{"orden": 1, "texto": "Nueva uno"}])
        self.assertEqual(response.data["competencias"], [integridad.pk])
        self.assertEqual(response.data["indicadores"], [{"orden": 1, "texto": "Meta"}])  # omitida: intacta
        descriptivo = DescriptivoPuesto.objects.get(pk=datos["id"])
        self.assertEqual(descriptivo.updated_by, self.gestor)

        vaciar = self.client.patch(url, {"funciones": []}, format="json")
        self.assertEqual(vaciar.data["funciones"], [])

    def test_no_se_puede_cambiar_la_posicion_de_un_descriptivo(self):
        datos = self._crear_por_api(self.create_posicion())
        response = self.client.patch(
            reverse("descriptivopuesto-detail", args=[datos["id"]]),
            {"posicion": str(self.create_posicion().pk)}, format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("posicion", response.data)

    def test_segundo_borrador_en_la_misma_posicion_se_rechaza_con_400(self):
        posicion = self.create_posicion()
        self._crear_por_api(posicion)
        response = self.client.post(reverse("descriptivopuesto-list"), self._payload(posicion), format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("posicion", response.data)

    def test_crear_borrador_precarga_desde_la_posicion(self):
        jefe = self.create_posicion(puesto=Puesto.objects.create(name="Supervisor de Prueba"))
        posicion = self.create_posicion(puesto=self.puesto, reports_to=jefe)
        self.client.force_authenticate(user=self.gestor)

        response = self.client.post(
            reverse("descriptivopuesto-crear-borrador"), {"posicion": str(posicion.pk)}, format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.assertEqual(response.data["nombre_puesto"], "Soldador de Prueba")
        self.assertEqual(response.data["reporta_a"], "Supervisor de Prueba")
        self.assertEqual(response.data["empresa"], self.company_node.name)
        self.assertEqual(response.data["version"], 1)
        self.assertEqual(DescriptivoPuesto.objects.get(pk=response.data["id"]).created_by, self.gestor)

        otra_vez = self.client.post(
            reverse("descriptivopuesto-crear-borrador"), {"posicion": str(posicion.pk)}, format="json",
        )
        self.assertEqual(otra_vez.status_code, status.HTTP_400_BAD_REQUEST)

    def test_crear_borrador_exige_una_posicion_que_exista(self):
        self.client.force_authenticate(user=self.gestor)
        sin_posicion = self.client.post(reverse("descriptivopuesto-crear-borrador"), {}, format="json")
        self.assertEqual(sin_posicion.status_code, status.HTTP_400_BAD_REQUEST)
        inexistente = self.client.post(
            reverse("descriptivopuesto-crear-borrador"),
            {"posicion": "00000000-0000-0000-0000-000000000000"}, format="json",
        )
        self.assertEqual(inexistente.status_code, status.HTTP_400_BAD_REQUEST)

    def test_congelar_deja_la_version_inmutable_para_edicion_y_borrado(self):
        posicion = self.create_posicion()
        liderazgo = CompetenciaConductual.objects.get(name="Liderazgo")
        datos = self._crear_por_api(
            posicion, proposito="Original", funciones=[{"texto": "Una"}], competencias=[liderazgo.pk],
        )
        url = reverse("descriptivopuesto-detail", args=[datos["id"]])

        congelado = self.client.post(reverse("descriptivopuesto-congelar", args=[datos["id"]]))
        self.assertEqual(congelado.status_code, status.HTTP_200_OK, congelado.data)
        self.assertTrue(congelado.data["esta_congelado"])
        self.assertIsNotNone(congelado.data["congelado_en"])

        self.assertEqual(
            self.client.post(reverse("descriptivopuesto-congelar", args=[datos["id"]])).status_code,
            status.HTTP_400_BAD_REQUEST,
        )
        edicion = self.client.patch(
            url, {"proposito": "Cambiado", "funciones": [{"texto": "Otra"}], "competencias": []}, format="json",
        )
        self.assertEqual(edicion.status_code, status.HTTP_400_BAD_REQUEST)
        solo_lista = self.client.patch(url, {"funciones": [{"texto": "Otra"}]}, format="json")
        self.assertEqual(solo_lista.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(self.client.delete(url).status_code, status.HTTP_400_BAD_REQUEST)

        intacto = self.client.get(url).data
        self.assertEqual(intacto["proposito"], "Original")
        self.assertEqual(intacto["funciones"], [{"orden": 1, "texto": "Una"}])
        self.assertEqual(intacto["competencias"], [liderazgo.pk])

    def test_borrar_un_borrador_es_borrado_logico(self):
        datos = self._crear_por_api(self.create_posicion())
        response = self.client.delete(reverse("descriptivopuesto-detail", args=[datos["id"]]))
        self.assertEqual(response.status_code, status.HTTP_204_NO_CONTENT)
        self.assertEqual(self.client.get(reverse("descriptivopuesto-list")).data["count"], 0)
        borrado = DescriptivoPuesto.all_objects.get(pk=datos["id"])
        self.assertTrue(borrado.is_deleted)
        self.assertEqual(borrado.deleted_by, self.gestor)

    def test_copiar_abre_un_borrador_nuevo_y_respeta_la_regla_de_uno_solo(self):
        datos = self._crear_por_api(self.create_posicion(), funciones=[{"texto": "Una"}])
        self.client.post(reverse("descriptivopuesto-congelar", args=[datos["id"]]))

        copia = self.client.post(reverse("descriptivopuesto-copiar", args=[datos["id"]]))
        self.assertEqual(copia.status_code, status.HTTP_201_CREATED, copia.data)
        self.assertEqual(copia.data["version"], 2)
        self.assertNotEqual(copia.data["id"], datos["id"])
        self.assertFalse(copia.data["esta_congelado"])
        self.assertEqual(copia.data["funciones"], [{"orden": 1, "texto": "Una"}])
        self.assertEqual(DescriptivoPuesto.objects.get(pk=copia.data["id"]).created_by, self.gestor)

        de_nuevo = self.client.post(reverse("descriptivopuesto-copiar", args=[datos["id"]]))
        self.assertEqual(de_nuevo.status_code, status.HTTP_400_BAD_REQUEST)

    def test_vigente_devuelve_la_congelada_mas_reciente(self):
        posicion = self.create_posicion()
        self.client.force_authenticate(user=self.colaborador)
        url = reverse("descriptivopuesto-vigente")
        self.assertEqual(self.client.get(url, {"posicion": str(posicion.pk)}).status_code, status.HTTP_404_NOT_FOUND)

        v1 = self.crear_descriptivo(posicion)
        # Un borrador no es vigente.
        self.assertEqual(self.client.get(url, {"posicion": str(posicion.pk)}).status_code, status.HTTP_404_NOT_FOUND)
        v1.congelar()
        self.assertEqual(self.client.get(url, {"posicion": str(posicion.pk)}).data["id"], str(v1.pk))
        v2 = copiar_version(v1)
        self.assertEqual(self.client.get(url, {"posicion": str(posicion.pk)}).data["id"], str(v1.pk))
        v2.congelar()
        self.assertEqual(self.client.get(url, {"posicion": str(posicion.pk)}).data["id"], str(v2.pk))

    def test_vigente_exige_una_posicion_valida(self):
        self.client.force_authenticate(user=self.colaborador)
        url = reverse("descriptivopuesto-vigente")
        self.assertEqual(self.client.get(url).status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(self.client.get(url, {"posicion": "no-es-un-id"}).status_code, status.HTTP_404_NOT_FOUND)
        self.assertEqual(
            self.client.get(url, {"posicion": "00000000-0000-0000-0000-000000000000"}).status_code,
            status.HTTP_404_NOT_FOUND,
        )

    def test_filtros_por_posicion_y_por_congelado(self):
        una, otra = self.create_posicion(), self.create_posicion()
        congelado = self.crear_descriptivo(una)
        congelado.congelar()
        borrador_de_una = copiar_version(congelado)
        borrador_de_otra = self.crear_descriptivo(otra)
        self.client.force_authenticate(user=self.colaborador)
        url = reverse("descriptivopuesto-list")

        def ids(**params):
            return {fila["id"] for fila in self.client.get(url, params).data["results"]}

        self.assertEqual(ids(posicion=str(una.pk)), {str(congelado.pk), str(borrador_de_una.pk)})
        self.assertEqual(ids(congelado="true"), {str(congelado.pk)})
        self.assertEqual(ids(congelado="false"), {str(borrador_de_una.pk), str(borrador_de_otra.pk)})
        self.assertEqual(ids(posicion=str(una.pk), congelado="false"), {str(borrador_de_una.pk)})

    def test_el_colaborador_no_ve_las_conformidades_pero_el_gestor_si(self):
        descriptivo = self.crear_descriptivo(self.create_posicion())
        descriptivo.congelar()
        ConformidadDescriptivo.objects.create(
            descriptivo=descriptivo, rol=RolConformidad.objects.get(name="Jefe inmediato"),
            fecha=date(2026, 3, 1), nombre_manual="Ana Jefa",
        )
        url = reverse("descriptivopuesto-detail", args=[descriptivo.pk])

        self.client.force_authenticate(user=self.colaborador)
        self.assertNotIn("conformidades", self.client.get(url).data)
        self.assertNotIn("conformidades", self.client.get(reverse("descriptivopuesto-list")).data["results"][0])

        self.client.force_authenticate(user=self.gestor)
        conformidades = self.client.get(url).data["conformidades"]
        self.assertEqual(len(conformidades), 1)
        self.assertEqual(conformidades[0]["nombre_manual"], "Ana Jefa")

    def test_cualquier_autenticado_puede_exportar_a_word_un_borrador_o_una_version(self):
        descriptivo = self.crear_descriptivo(self.create_posicion(), nombre_puesto="Soldador")
        self.client.force_authenticate(user=self.colaborador)
        url = reverse("descriptivopuesto-exportar-word", args=[descriptivo.pk])

        borrador = self.client.get(url)
        self.assertEqual(borrador.status_code, status.HTTP_200_OK)
        self.assertEqual(
            borrador["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn("attachment", borrador["Content-Disposition"])
        self.assertIn("FO-C0-CH-04", borrador["Content-Disposition"])
        contenido = b"".join(borrador.streaming_content)
        with ZipFile(io.BytesIO(contenido)) as paquete:
            self.assertIn("Soldador", paquete.read("word/document.xml").decode("utf-8"))

        descriptivo.congelar()
        self.assertEqual(self.client.get(url).status_code, status.HTTP_200_OK)

    def test_exportar_a_word_exige_autenticacion_y_un_descriptivo_que_exista(self):
        descriptivo = self.crear_descriptivo(self.create_posicion())
        url = reverse("descriptivopuesto-exportar-word", args=[descriptivo.pk])
        self.assertEqual(self.client.get(url).status_code, status.HTTP_401_UNAUTHORIZED)
        self.client.force_authenticate(user=self.colaborador)
        inexistente = reverse("descriptivopuesto-exportar-word", args=["00000000-0000-0000-0000-000000000000"])
        self.assertEqual(self.client.get(inexistente).status_code, status.HTTP_404_NOT_FOUND)


class ConformidadDescriptivoAPITests(DescriptivoPuestoTestMixin, APITestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        user_model = get_user_model()
        cls.gestor = user_model.objects.create_user(
            username="conformidad-gestor", email="conformidad-gestor@example.com",
            password="strong-test-password", role=UserRole.objects.get(code="capital-humano"),
        )
        cls.colaborador = user_model.objects.create_user(
            username="conformidad-colaborador", email="conformidad-colaborador@example.com",
            password="strong-test-password", role=UserRole.objects.get(code="colaborador"),
        )
        cls.rol_jefe = RolConformidad.objects.get(name="Jefe inmediato")
        cls.rol_colaborador = RolConformidad.objects.get(name="Colaborador")

    def _congelado(self):
        descriptivo = self.crear_descriptivo(self.create_posicion())
        descriptivo.congelar()
        return descriptivo

    def test_un_colaborador_no_puede_leer_ni_crear_conformidades(self):
        self.client.force_authenticate(user=self.colaborador)
        self.assertEqual(self.client.get(reverse("conformidaddescriptivo-list")).status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(
            self.client.post(reverse("conformidaddescriptivo-list"), {}, format="json").status_code,
            status.HTTP_403_FORBIDDEN,
        )

    def test_gestor_registra_una_conformidad_sobre_una_version_congelada(self):
        descriptivo = self._congelado()
        persona = self.crear_persona()
        self.client.force_authenticate(user=self.gestor)
        response = self.client.post(
            reverse("conformidaddescriptivo-list"),
            {
                "descriptivo": str(descriptivo.pk), "rol": str(self.rol_colaborador.pk),
                "persona": str(persona.pk), "fecha": "2026-03-01",
            },
            format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED, response.data)
        self.assertEqual(ConformidadDescriptivo.objects.get(pk=response.data["id"]).created_by, self.gestor)

    def test_la_conformidad_trae_el_nombre_de_la_persona_sin_espacio_doble(self):
        descriptivo = self._congelado()
        persona = self.crear_persona("Luis")  # sin apellido materno
        self.client.force_authenticate(user=self.gestor)
        creada = self.client.post(
            reverse("conformidaddescriptivo-list"),
            {
                "descriptivo": str(descriptivo.pk), "rol": str(self.rol_colaborador.pk),
                "persona": str(persona.pk), "fecha": "2026-03-01",
            },
            format="json",
        )
        self.assertEqual(creada.data["persona_nombre"], "Prueba Luis")
        # Y el detalle del descriptivo la trae sin una consulta por conformidad.
        detalle = self.client.get(reverse("descriptivopuesto-detail", args=[descriptivo.pk])).data
        self.assertEqual(detalle["conformidades"][0]["persona_nombre"], "Prueba Luis")

    def test_no_se_registra_conformidad_sobre_un_borrador(self):
        borrador = self.crear_descriptivo(self.create_posicion())
        self.client.force_authenticate(user=self.gestor)
        response = self.client.post(
            reverse("conformidaddescriptivo-list"),
            {"descriptivo": str(borrador.pk), "rol": str(self.rol_jefe.pk)}, format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("descriptivo", response.data)

    def test_el_colaborador_exige_persona_tambien_por_la_api(self):
        descriptivo = self._congelado()
        self.client.force_authenticate(user=self.gestor)
        response = self.client.post(
            reverse("conformidaddescriptivo-list"),
            {"descriptivo": str(descriptivo.pk), "rol": str(self.rol_colaborador.pk)}, format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        # El mensaje del modelo, no un genérico "campo requerido" del serializer.
        self.assertIn("exige indicar la persona", str(response.data["persona"]))

    def _registrar(self, descriptivo, rol, **extra):
        return self.client.post(
            reverse("conformidaddescriptivo-list"),
            {"descriptivo": str(descriptivo.pk), "rol": str(rol.pk), **extra}, format="json",
        )

    def test_jefe_y_capital_humano_se_registran_sin_mandar_persona(self):
        descriptivo = self._congelado()
        self.client.force_authenticate(user=self.gestor)
        jefe = self._registrar(descriptivo, self.rol_jefe, fecha="2026-03-01", nombre_manual="Ana Jefa")
        self.assertEqual(jefe.status_code, status.HTTP_201_CREATED, jefe.data)
        self.assertIsNone(jefe.data["persona"])
        rh = self._registrar(
            descriptivo, RolConformidad.objects.get(name="Capital Humano"), fecha="2026-03-02", nombre_manual="Rosa",
        )
        self.assertEqual(rh.status_code, status.HTTP_201_CREATED, rh.data)

    def test_dos_personas_distintas_dan_su_conformidad_como_colaborador_por_la_api(self):
        descriptivo = self._congelado()
        self.client.force_authenticate(user=self.gestor)
        primera = self._registrar(
            descriptivo, self.rol_colaborador, persona=str(self.crear_persona("Luis").pk), fecha="2026-03-01",
        )
        segunda = self._registrar(
            descriptivo, self.rol_colaborador, persona=str(self.crear_persona("Mara").pk), fecha="2027-03-01",
        )
        self.assertEqual(primera.status_code, status.HTTP_201_CREATED, primera.data)
        self.assertEqual(segunda.status_code, status.HTTP_201_CREATED, segunda.data)

    def test_la_misma_persona_o_el_mismo_rol_sin_persona_no_se_repiten_por_la_api(self):
        descriptivo = self._congelado()
        persona = self.crear_persona()
        self.client.force_authenticate(user=self.gestor)
        self.assertEqual(
            self._registrar(descriptivo, self.rol_colaborador, persona=str(persona.pk), fecha="2026-03-01").status_code,
            status.HTTP_201_CREATED,
        )
        self.assertEqual(
            self._registrar(descriptivo, self.rol_colaborador, persona=str(persona.pk)).status_code,
            status.HTTP_400_BAD_REQUEST,
        )
        self.assertEqual(
            self._registrar(descriptivo, self.rol_jefe, fecha="2026-03-01", nombre_manual="Ana").status_code,
            status.HTTP_201_CREATED,
        )
        self.assertEqual(self._registrar(descriptivo, self.rol_jefe).status_code, status.HTTP_400_BAD_REQUEST)

    def test_una_conformidad_no_se_puede_mover_a_otra_version(self):
        original = self._congelado()
        otra = self._congelado()
        conformidad = ConformidadDescriptivo.objects.create(
            descriptivo=original, rol=self.rol_jefe, fecha=date(2026, 3, 1), nombre_manual="Ana",
        )
        self.client.force_authenticate(user=self.gestor)
        response = self.client.patch(
            reverse("conformidaddescriptivo-detail", args=[conformidad.pk]),
            {"descriptivo": str(otra.pk)}, format="json",
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        conformidad.refresh_from_db()
        self.assertEqual(conformidad.descriptivo_id, original.pk)

    def test_listado_filtra_por_descriptivo_y_borrar_es_logico(self):
        uno, otro = self._congelado(), self._congelado()
        propia = ConformidadDescriptivo.objects.create(
            descriptivo=uno, rol=self.rol_jefe, fecha=date(2026, 3, 1), nombre_manual="Ana",
        )
        ConformidadDescriptivo.objects.create(
            descriptivo=otro, rol=self.rol_jefe, fecha=date(2026, 3, 1), nombre_manual="Beto",
        )
        self.client.force_authenticate(user=self.gestor)

        filtradas = self.client.get(reverse("conformidaddescriptivo-list"), {"descriptivo": str(uno.pk)})
        self.assertEqual([fila["id"] for fila in filtradas.data["results"]], [str(propia.pk)])

        borrada = self.client.delete(reverse("conformidaddescriptivo-detail", args=[propia.pk]))
        self.assertEqual(borrada.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(ConformidadDescriptivo.objects.filter(pk=propia.pk).exists())
        self.assertTrue(ConformidadDescriptivo.all_objects.get(pk=propia.pk).is_deleted)


class AdminDescargasOficialesTests(DescriptivoPuestoTestMixin, TestCase):
    """
    Acciones del admin para sacar los archivos oficiales sin frontend ni
    token: los mismos que entregan exportar-excel / exportar-word.
    """

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.admin_user = get_user_model().objects.create_superuser(
            username="admin-descargas", email="admin-descargas@example.com", password="test-pass",
        )

    def setUp(self):
        self.client.force_login(self.admin_user)

    def _accion(self, nombre_url, accion, seleccion):
        return self.client.post(
            reverse(nombre_url),
            {"action": accion, "_selected_action": [str(pk) for pk in seleccion]},
        )

    def test_una_requisicion_baja_su_excel_oficial(self):
        requisicion = self.create_requisicion(area_solicitante="Mantenimiento")
        response = self._accion(
            "admin:recruitment_requisicion_changelist", "descargar_excel_oficial", [requisicion.pk],
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(f"_{requisicion.pk}.xlsx", response["Content-Disposition"])
        ws = openpyxl.load_workbook(io.BytesIO(b"".join(response.streaming_content))).active
        self.assertEqual(ws["M4"].value, "Mantenimiento")

    def test_varias_requisiciones_bajan_en_un_zip(self):
        una = self.create_requisicion()
        otra = self.create_requisicion(tipo=self.tipo_nueva, justificacion="Crecimiento del área.")
        response = self._accion(
            "admin:recruitment_requisicion_changelist", "descargar_excel_oficial", [una.pk, otra.pk],
        )
        self.assertEqual(response["Content-Type"], "application/zip")
        with ZipFile(io.BytesIO(b"".join(response.streaming_content))) as paquete:
            nombres = sorted(paquete.namelist())
            self.assertEqual(len(nombres), 2)
            self.assertTrue(any("FO-C0-CH-08" in n for n in nombres))  # Reemplazo
            self.assertTrue(any("FO-C0-CH-01" in n for n in nombres))  # Nueva Posicion
            self.assertIsNone(paquete.testzip())

    def test_un_tipo_sin_plantilla_avisa_y_no_truena(self):
        raro = TipoRequisicion.objects.create(name="Tipo sin plantilla")
        requisicion = self.create_requisicion(tipo=raro)
        response = self._accion(
            "admin:recruitment_requisicion_changelist", "descargar_excel_oficial", [requisicion.pk],
        )
        self.assertEqual(response.status_code, 302)  # vuelve a la lista
        avisos = [str(m) for m in get_messages(response.wsgi_request)]
        self.assertTrue(any("No hay plantilla oficial" in aviso for aviso in avisos), avisos)

    def test_un_descriptivo_baja_su_word_oficial(self):
        descriptivo = self.crear_descriptivo(self.create_posicion(), nombre_puesto="Soldador")
        response = self._accion(
            "admin:recruitment_descriptivopuesto_changelist", "descargar_word_oficial", [descriptivo.pk],
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"], "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        with ZipFile(io.BytesIO(b"".join(response.streaming_content))) as paquete:
            self.assertIn("Soldador", paquete.read("word/document.xml").decode("utf-8"))

    def test_varios_descriptivos_bajan_en_un_zip(self):
        uno = self.crear_descriptivo(self.create_posicion(), nombre_puesto="Soldador")
        otro = self.crear_descriptivo(self.create_posicion(), nombre_puesto="Pintor")
        response = self._accion(
            "admin:recruitment_descriptivopuesto_changelist", "descargar_word_oficial", [uno.pk, otro.pk],
        )
        self.assertEqual(response["Content-Type"], "application/zip")
        with ZipFile(io.BytesIO(b"".join(response.streaming_content))) as paquete:
            self.assertEqual(len(paquete.namelist()), 2)
            self.assertTrue(all(n.endswith(".docx") for n in paquete.namelist()))

    def test_las_acciones_aparecen_en_el_admin(self):
        # Con la lista vacía el admin no dibuja la barra de acciones.
        self.create_requisicion()
        self.crear_descriptivo(self.create_posicion())
        for nombre_url, accion in (
            ("admin:recruitment_requisicion_changelist", "descargar_excel_oficial"),
            ("admin:recruitment_descriptivopuesto_changelist", "descargar_word_oficial"),
        ):
            with self.subTest(accion=accion):
                self.assertContains(self.client.get(reverse(nombre_url)), f'value="{accion}"')


class AdminListasPorTipoTests(RecruitmentTestDataMixin, TestCase):
    """
    El admin muestra Reemplazo y Nueva Posición como dos listas separadas
    (modelos proxy sobre la misma tabla), cada una con su formulario oficial.
    """

    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        cls.admin_user = get_user_model().objects.create_superuser(
            username="admin-listas", email="admin-listas@example.com", password="test-pass",
        )
        cls.reemplazo = cls.create_requisicion(tipo=cls.tipo_reemplazo, area_solicitante="De reemplazo")
        cls.nueva = cls.create_requisicion(
            tipo=cls.tipo_nueva, justificacion="Crecimiento.", area_solicitante="De nueva posición",
        )

    def setUp(self):
        self.client.force_login(self.admin_user)

    def _ids_de_la_lista(self, nombre_url):
        respuesta = self.client.get(reverse(nombre_url))
        self.assertEqual(respuesta.status_code, 200)
        return {str(o.pk) for o in respuesta.context["cl"].queryset}

    def _datos_alta(self, posicion, **extra):
        return {
            "posicion": str(posicion.pk), "estado": str(self.estado_borrador.pk), "fecha_solicitud": "2026-02-01",
            "aprobaciones-TOTAL_FORMS": "0", "aprobaciones-INITIAL_FORMS": "0",
            "aprobaciones-MIN_NUM_FORMS": "0", "aprobaciones-MAX_NUM_FORMS": "1000",
            **extra,
        }

    def test_cada_lista_muestra_solo_su_tipo_y_la_general_muestra_todas(self):
        self.assertEqual(self._ids_de_la_lista("admin:recruitment_requisicionreemplazo_changelist"), {str(self.reemplazo.pk)})
        self.assertEqual(self._ids_de_la_lista("admin:recruitment_requisicionnuevaposicion_changelist"), {str(self.nueva.pk)})
        self.assertEqual(
            self._ids_de_la_lista("admin:recruitment_requisicion_changelist"), {str(self.reemplazo.pk), str(self.nueva.pk)},
        )

    def test_las_listas_tambien_muestran_las_borradas_para_poder_auditarlas(self):
        self.reemplazo.deleted_by = self.admin_user
        self.reemplazo.delete()
        self.assertIn(str(self.reemplazo.pk), self._ids_de_la_lista("admin:recruitment_requisicionreemplazo_changelist"))

    def test_un_registro_no_se_alcanza_desde_la_lista_del_otro_tipo(self):
        url = reverse("admin:recruitment_requisicionreemplazo_change", args=[self.nueva.pk])
        respuesta = self.client.get(url)
        self.assertEqual(respuesta.status_code, 302)  # el admin lo trata como "no existe"
        propia = reverse("admin:recruitment_requisicionreemplazo_change", args=[self.reemplazo.pk])
        self.assertEqual(self.client.get(propia).status_code, 200)

    def test_alta_en_la_lista_de_reemplazo_pone_el_tipo_sola_y_no_pide_justificacion(self):
        posicion = self.create_posicion()
        respuesta = self.client.post(
            reverse("admin:recruitment_requisicionreemplazo_add"), self._datos_alta(posicion),
        )
        self.assertEqual(respuesta.status_code, 302, getattr(respuesta, "context", None) and respuesta.context["adminform"].form.errors)
        creada = Requisicion.objects.get(posicion=posicion)
        self.assertEqual(creada.tipo, self.tipo_reemplazo)
        self.assertEqual(creada.created_by, self.admin_user)

    def test_alta_en_la_lista_de_nueva_posicion_exige_justificacion(self):
        posicion = self.create_posicion()
        url = reverse("admin:recruitment_requisicionnuevaposicion_add")

        sin = self.client.post(url, self._datos_alta(posicion))
        self.assertEqual(sin.status_code, 200)  # vuelve al formulario con el error
        self.assertIn("justificacion", sin.context["adminform"].form.errors)
        self.assertFalse(Requisicion.objects.filter(posicion=posicion).exists())

        con = self.client.post(url, self._datos_alta(posicion, justificacion="Se abre un segundo turno."))
        self.assertEqual(con.status_code, 302)
        creada = Requisicion.objects.get(posicion=posicion)
        self.assertEqual(creada.tipo, self.tipo_nueva)
        self.assertEqual(creada.justificacion, "Se abre un segundo turno.")

    def test_el_tipo_no_se_puede_cambiar_desde_la_lista_de_un_tipo(self):
        url = reverse("admin:recruitment_requisicionreemplazo_change", args=[self.reemplazo.pk])
        datos = self._datos_alta(self.reemplazo.posicion, tipo=str(self.tipo_nueva.pk), area_solicitante="Editada")
        datos["aprobaciones-INITIAL_FORMS"] = "0"
        respuesta = self.client.post(url, datos)
        self.assertEqual(respuesta.status_code, 302)
        self.reemplazo.refresh_from_db()
        self.assertEqual(self.reemplazo.area_solicitante, "Editada")
        self.assertEqual(self.reemplazo.tipo, self.tipo_reemplazo)  # el tipo que llegó en el POST se ignora

    def test_cada_lista_baja_su_formulario_oficial(self):
        for nombre_url, requisicion, codigo in (
            ("admin:recruitment_requisicionreemplazo_changelist", self.reemplazo, "FO-C0-CH-08"),
            ("admin:recruitment_requisicionnuevaposicion_changelist", self.nueva, "FO-C0-CH-01"),
        ):
            with self.subTest(formulario=codigo):
                respuesta = self.client.post(
                    reverse(nombre_url),
                    {"action": "descargar_excel_oficial", "_selected_action": [str(requisicion.pk)]},
                )
                self.assertEqual(respuesta.status_code, 200)
                self.assertIn(codigo, respuesta["Content-Disposition"])

    def test_sin_el_tipo_en_el_catalogo_no_se_ofrece_crear_en_esa_lista(self):
        # Una base nueva sin el catálogo de tipos (la sábana aún no se importa)
        # no debe romper la pantalla de alta: simplemente no se permite crear.
        from apps.recruitment.admin import RequisicionNuevaPosicionAdmin

        with mock.patch.object(RequisicionNuevaPosicionAdmin, "codigo_tipo", "tipo-que-no-existe"):
            respuesta = self.client.get(reverse("admin:recruitment_requisicionnuevaposicion_add"))
        self.assertEqual(respuesta.status_code, 403)
        # Con el tipo presente, la misma pantalla sí abre.
        self.assertEqual(self.client.get(reverse("admin:recruitment_requisicionnuevaposicion_add")).status_code, 200)


class PosicionesElegiblesAPITests(DescriptivoPuestoTestMixin, APITestCase):
    @classmethod
    def setUpTestData(cls):
        super().setUpTestData()
        user_model = get_user_model()
        for atributo, rol in (
            ("colaborador", "colaborador"), ("otro", "colaborador"),
            ("gestor", "capital-humano"), ("administrador", "admin"),
        ):
            setattr(cls, atributo, user_model.objects.create_user(
                username=f"selector-{atributo}", email=f"selector-{atributo}@example.com",
                password="strong-test-password", role=UserRole.objects.get(code=rol),
            ))
        cls.puesto = Puesto.objects.create(name="Operador de Prueba")
        cls.area = Area.objects.create(code="SELECTOR", name="Producción")
        cls.estatus = {"vacante-activa": cls.estatus_vacante}
        for code in (
            "vacante-pendiente-de-confirmacion", "vacante-suspendida", "vacante-eliminada",
            "colaborador-activo", "colaborador-baja", "trainee-activo", "trainee-baja",
        ):
            cls.estatus[code] = EstatusPosicion.objects.create(code=code, name=code)

    def setUp(self):
        self.client.force_authenticate(user=self.colaborador)
        self.url = reverse("posiciones-elegibles-list")

    def _listar(self, para="requisicion", **params):
        response = self.client.get(self.url, {"para": para, **params})
        self.assertEqual(response.status_code, status.HTTP_200_OK, response.data)
        return response.data

    def test_para_es_obligatorio_y_solo_admite_los_dos_tramites(self):
        for params in ({}, {"para": ""}, {"para": "otro"}):
            with self.subTest(params=params):
                response = self.client.get(self.url, params)
                self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
                self.assertEqual(str(response.data["para"]), "Indica para=requisicion o para=descriptivo.")

    def test_un_anonimo_no_puede_consultar(self):
        self.client.force_authenticate(user=None)
        response = self.client.get(self.url, {"para": "requisicion"})
        self.assertIn(response.status_code, (status.HTTP_401_UNAUTHORIZED, status.HTTP_403_FORBIDDEN))

    def test_solo_admite_get(self):
        for metodo in ("post", "put", "patch", "delete"):
            with self.subTest(metodo=metodo):
                response = getattr(self.client, metodo)(self.url, {}, format="json")
                self.assertEqual(response.status_code, status.HTTP_405_METHOD_NOT_ALLOWED)

    def test_ordena_por_code_luego_puesto_y_pk_sin_ocultar_posiciones(self):
        puesto_a = Puesto.objects.create(name="Auxiliar de Prueba")
        posiciones = []
        for code in reversed(tuple(self.estatus)):
            # El nombre no determina ni la prioridad ni si está ocupada.
            estatus = self.estatus[code]
            estatus.name = "Nombre de catálogo cambiado"
            estatus.save()
            posiciones.append(self.create_posicion(estatus=estatus, puesto=self.puesto))
        posiciones.extend(self.create_posicion(puesto=puesto_a) for _ in range(2))
        prioridad = {
            "vacante-activa": 0, "vacante-pendiente-de-confirmacion": 1,
            "vacante-suspendida": 2, "vacante-eliminada": 4,
        }
        esperadas = sorted(posiciones, key=lambda p: (
            prioridad.get(p.estatus.code, 3), p.puesto.name, p.pk,
        ))
        for para in ("requisicion", "descriptivo"):
            with self.subTest(para=para):
                datos = self._listar(para)["results"]
                self.assertEqual([dato["id"] for dato in datos], [str(p.pk) for p in esperadas])
                for dato in datos:
                    self.assertEqual(dato["ocupada"], dato["estatus_code"].startswith(("colaborador-", "trainee-")))

    def test_ordena_sin_tramite_primero_dentro_de_cada_prioridad_para_ambos_selectores(self):
        puesto_a = Puesto.objects.create(name="Auxiliar de Prueba")
        esperadas = {"requisicion": [], "descriptivo": []}
        for code in (
            "vacante-activa", "vacante-pendiente-de-confirmacion", "vacante-suspendida",
            "colaborador-activo", "vacante-eliminada",
        ):
            con_requisicion = [self.create_posicion(estatus=self.estatus[code], puesto=puesto_a) for _ in range(2)]
            con_descriptivo = [self.create_posicion(estatus=self.estatus[code], puesto=self.puesto) for _ in range(2)]
            for posicion in con_requisicion:
                # Un trámite ajeno también cuenta aunque no se pueda leer su id.
                self.create_requisicion(posicion=posicion, created_by=self.otro)
            for posicion in con_descriptivo:
                self.crear_descriptivo(posicion)
            con_requisicion.sort(key=lambda p: p.pk)
            con_descriptivo.sort(key=lambda p: p.pk)
            esperadas["requisicion"].extend(con_descriptivo + con_requisicion)
            esperadas["descriptivo"].extend(con_requisicion + con_descriptivo)
        for para in ("requisicion", "descriptivo"):
            with self.subTest(para=para):
                datos = self._listar(para)["results"]
                self.assertEqual([dato["id"] for dato in datos], [str(p.pk) for p in esperadas[para]])

    def test_devuelve_los_campos_y_la_etiqueta_existente(self):
        posicion = self.create_posicion(puesto=self.puesto, area=self.area)
        self.assertEqual(self._listar()["results"], [{
            "id": str(posicion.pk), "etiqueta": posicion.etiqueta, "puesto": self.puesto.name,
            "unidad": self.company_node.name, "area": self.area.name,
            "estatus": self.estatus_vacante.name, "estatus_code": self.estatus_vacante.code,
            "ocupada": False, "tramite_abierto": None,
        }])

    def test_conserva_posiciones_sin_puesto_ni_area(self):
        posicion = self.create_posicion()
        dato = self._listar()["results"][0]
        self.assertEqual(dato["etiqueta"], posicion.etiqueta)
        self.assertIsNone(dato["puesto"])
        self.assertIsNone(dato["area"])

    def test_marca_requisicion_abierta_solo_en_su_selector(self):
        posicion = self.create_posicion(puesto=self.puesto)
        requisicion = self.create_requisicion(posicion=posicion, created_by=self.colaborador)
        self.assertEqual(self._listar()["results"][0]["tramite_abierto"], {
            "tipo": "requisicion", "id": str(requisicion.pk), "estado": self.estado_borrador.name,
        })
        self.assertIsNone(self._listar("descriptivo")["results"][0]["tramite_abierto"])

    def test_requisicion_terminal_no_cuenta_como_abierta(self):
        posicion = self.create_posicion()
        self.create_requisicion(posicion=posicion, estado=self.estado_cubierta, created_by=self.colaborador)
        self.assertIsNone(self._listar()["results"][0]["tramite_abierto"])

    def test_requisicion_borrada_no_cuenta_como_abierta(self):
        posicion = self.create_posicion()
        self.create_requisicion(posicion=posicion, created_by=self.colaborador).delete()
        self.assertIsNone(self._listar()["results"][0]["tramite_abierto"])

    def test_privacidad_de_requisicion_para_ajeno_dueno_y_gestion(self):
        posicion = self.create_posicion()
        requisicion = self.create_requisicion(posicion=posicion, created_by=self.colaborador)
        for usuario in (self.otro, self.colaborador, self.gestor, self.administrador):
            with self.subTest(usuario=usuario.username):
                self.client.force_authenticate(user=usuario)
                puede_leer = usuario != self.otro
                self.assertEqual(self._listar()["results"][0]["tramite_abierto"], {
                    "tipo": "requisicion", "id": str(requisicion.pk) if puede_leer else None,
                    "estado": self.estado_borrador.name if puede_leer else None,
                })

    def test_requisicion_sin_dueno_tambien_oculta_id_y_estado_al_colaborador(self):
        self.create_requisicion(posicion=self.create_posicion())
        self.assertEqual(self._listar()["results"][0]["tramite_abierto"], {
            "tipo": "requisicion", "id": None, "estado": None,
        })

    def test_marca_borrador_de_descriptivo_para_cualquier_autenticado(self):
        posicion = self.create_posicion()
        borrador = self.crear_descriptivo(posicion, created_by=self.gestor)
        for usuario in (self.colaborador, self.gestor):
            with self.subTest(usuario=usuario.username):
                self.client.force_authenticate(user=usuario)
                self.assertEqual(self._listar("descriptivo")["results"][0]["tramite_abierto"], {
                    "tipo": "descriptivo", "id": str(borrador.pk), "estado": "Borrador",
                })
                self.assertIsNone(self._listar()["results"][0]["tramite_abierto"])

    def test_descriptivo_congelado_no_cuenta_y_el_nuevo_borrador_si(self):
        posicion = self.create_posicion()
        congelado = self.crear_descriptivo(posicion)
        congelado.congelar()
        self.assertIsNone(self._listar("descriptivo")["results"][0]["tramite_abierto"])
        borrador = self.crear_descriptivo(posicion)
        self.assertEqual(self._listar("descriptivo")["results"][0]["tramite_abierto"]["id"], str(borrador.pk))

    def test_descriptivo_borrado_no_cuenta_como_abierto(self):
        posicion = self.create_posicion()
        self.crear_descriptivo(posicion).delete()
        self.assertIsNone(self._listar("descriptivo")["results"][0]["tramite_abierto"])

    def test_busca_sin_acentos_en_los_mismos_campos_de_posiciones(self):
        puesto = Puesto.objects.create(name="Operación")
        jefe = self.create_posicion(puesto=Puesto.objects.create(name="Supervisión"))
        self.company_node.name = "Organización"
        self.company_node.save()
        self.estatus_vacante.name = "Vacante con Confirmación"
        self.estatus_vacante.save()
        posicion = self.create_posicion(puesto=puesto, area=self.area, reports_to=jefe)
        for para in ("requisicion", "descriptivo"):
            for search in ("OPERACION", "produccion", "organizacion", "confirmacion", "supervision"):
                with self.subTest(para=para, search=search):
                    datos = self._listar(para, search=search)["results"]
                    self.assertIn(str(posicion.pk), [dato["id"] for dato in datos])
            self.assertEqual(self._listar(para, search="inexistente")["results"], [])
        puesto.name = "Operacion"
        puesto.save()
        self.assertEqual(self._listar(search="operación")["results"][0]["id"], str(posicion.pk))

    def test_busca_etiquetas_con_separadores_y_parentesis(self):
        self.company_node.name = "GPA"
        self.company_node.save()
        posicion = self.create_posicion(
            puesto=Puesto.objects.create(name="Operador de Soldadura"), area=self.area,
        )
        self.create_posicion(puesto=Puesto.objects.create(name="Auxiliar"))
        posicion = Posicion.objects.get(pk=posicion.pk)
        for para in ("requisicion", "descriptivo"):
            for search in (
                "Operador de Soldadura — GPA", "Operador de Soldadura (Produccion)",
                "Operador - ( Produccion )", posicion.etiqueta,
            ):
                with self.subTest(para=para, search=search):
                    datos = self._listar(para, search=search)["results"]
                    self.assertEqual([dato["id"] for dato in datos], [str(posicion.pk)])
            self.assertEqual(self._listar(para, search="Operador — inexistente")["results"], [])

    def test_paginacion_por_defecto_personalizada_y_maximo(self):
        Posicion.objects.bulk_create([
            Posicion(organization_node=self.company_node, estatus=self.estatus_vacante, puesto=self.puesto)
            for _ in range(205)
        ])
        datos = self._listar()
        self.assertEqual(datos["count"], 205)
        self.assertEqual(len(datos["results"]), 25)
        self.assertIsNotNone(datos["next"])
        primera = self._listar(page_size=8)
        segunda = self._listar(page_size=8, page=2)
        self.assertEqual(len(primera["results"]), 8)
        self.assertEqual(len(segunda["results"]), 8)
        self.assertTrue(
            {dato["id"] for dato in primera["results"]}.isdisjoint(dato["id"] for dato in segunda["results"])
        )
        self.assertEqual(len(self._listar(page_size=999)["results"]), 200)

    def test_numero_de_consultas_constante_con_8_y_40_posiciones(self):
        # Aislamos las consultas del endpoint de la carga del rol del usuario.
        self.colaborador.role
        for cantidad in (8, 40):
            for _ in range(cantidad - Posicion.objects.count()):
                posicion = self.create_posicion(puesto=self.puesto, area=self.area)
                self.create_requisicion(posicion=posicion, created_by=self.colaborador)
                self.crear_descriptivo(posicion, created_by=self.gestor)
            for para in ("requisicion", "descriptivo"):
                with self.subTest(cantidad=cantidad, para=para):
                    # Un COUNT para paginar y un SELECT con relaciones y subconsultas.
                    with self.assertNumQueries(2):
                        datos = self._listar(para, page_size=200)
                    self.assertEqual(len(datos["results"]), cantidad)
                    self.assertTrue(all(dato["tramite_abierto"]["id"] for dato in datos["results"]))

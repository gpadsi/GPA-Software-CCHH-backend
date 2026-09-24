# apps/imports/management/commands/import_configuracion.py
# Lee la hoja "Configuración" de la sábana de GPA: es el catálogo que GPA
# usa para sus propios menús desplegables, no texto libre — por eso los
# valores nuevos (Escolaridad, Estado Civil) se siembran directo, y los que
# ya teníamos confirmados (Estatus, Alcance, Tipo de Requisición, Género)
# solo se comparan y se reportan si hay diferencia, nunca se sobrescriben
# solos.
import openpyxl
from django.core.management.base import BaseCommand, CommandError

from apps.imports.normalize import normalize_text
from apps.persons.models import Escolaridad, EstadoCivil, Genero
from apps.positions.models import AlcanceDePosicion, EstatusPosicion, TipoRequisicion

# label tal como aparece literalmente en la hoja -> (modo, modelo)
# "seed": catálogo nuevo, se crea lo que falte.
# "check": catálogo ya confirmado, solo se reporta si no coincide.
HEADER_TARGETS = {
    "Escolaridad": ("seed", Escolaridad),
    "Estado Civil": ("seed", EstadoCivil),
    "Alcance": ("check", AlcanceDePosicion),
    "Tipo de Requisición": ("check", TipoRequisicion),
    "Estatus": ("check", EstatusPosicion),
    "Genero": ("check", Genero),
}


class Command(BaseCommand):
    help = (
        "Importa la hoja 'Configuración' de la sábana de GPA: siembra "
        "Escolaridad/Estado Civil (nuevos) y reporta diferencias contra "
        "Estatus/Alcance/Tipo de Requisición/Género (ya confirmados)."
    )

    def add_arguments(self, parser):
        parser.add_argument("path")
        parser.add_argument("--sheet", default="Configuración")

    def handle(self, *args, **options):
        path = options["path"]
        sheet_name = options["sheet"]

        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise CommandError(f"No se pudo abrir {path}: {exc}")

        if sheet_name not in wb.sheetnames:
            raise CommandError(f"La hoja «{sheet_name}» no existe en {path}. Hojas disponibles: {wb.sheetnames}")

        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        blocks = self._extract_blocks(rows, HEADER_TARGETS.keys())

        for label, (mode, model) in HEADER_TARGETS.items():
            values = blocks.get(label)
            if not values:
                self.stdout.write(self.style.WARNING(f"No encontré el bloque «{label}» en la hoja — nada que hacer."))
                continue
            if mode == "seed":
                self._seed(model, values)
            else:
                self._check(model, values, label)

    def _extract_blocks(self, rows, labels):
        """
        Encuentra cada label como celda de texto exacta, y lee hacia abajo en
        esa misma columna hasta la primera celda vacía — así un bloque
        distinto más abajo en la misma columna nunca se mezcla con el
        anterior (siempre hay al menos una fila en blanco entre bloques).
        """
        labels = set(labels)
        found_at = {}
        for r_idx, row in enumerate(rows):
            for c_idx, cell in enumerate(row):
                if isinstance(cell, str) and cell.strip() in labels and cell.strip() not in found_at:
                    found_at[cell.strip()] = (r_idx, c_idx)

        blocks = {}
        for label, (r_idx, c_idx) in found_at.items():
            values = []
            r = r_idx + 1
            while r < len(rows):
                row = rows[r]
                cell = row[c_idx] if c_idx < len(row) else None
                if cell is None or (isinstance(cell, str) and not cell.strip()):
                    break
                values.append(str(cell).strip())
                r += 1
            blocks[label] = values
        return blocks

    def _seed(self, model, values):
        created = 0
        for name in values:
            _, was_created = model.objects.get_or_create(name=name)
            created += int(was_created)
        self.stdout.write(self.style.SUCCESS(
            f"{model.__name__}: {created} creados, {len(values) - created} ya existían "
            f"(total en archivo: {len(values)})."
        ))

    def _check(self, model, values, label):
        existing = {normalize_text(o.name) for o in model.objects.all()}
        from_file = {normalize_text(v) for v in values}
        missing_in_db = from_file - existing
        missing_in_file = existing - from_file

        if not missing_in_db and not missing_in_file:
            self.stdout.write(self.style.SUCCESS(f"{label}: coincide exactamente con lo ya sembrado ({len(existing)} valores)."))
            return

        if missing_in_db:
            self.stdout.write(self.style.WARNING(f"{label}: en el archivo pero NO sembrados todavía: {sorted(missing_in_db)}"))
        if missing_in_file:
            self.stdout.write(self.style.WARNING(f"{label}: sembrados pero NO aparecen en este archivo: {sorted(missing_in_file)}"))

# apps/imports/management/commands/import_colaboradores.py
# Importa "Lista Colaboradores" de GPA: crea/actualiza Persona + Empleado
# (identidad y datos de contacto). A PROPÓSITO no crea Posicion ni Contrato
# aquí — este archivo no trae estructura organizacional/física suficiente
# para eso; Posicion/Contrato los crea el import de la sábana de Posiciones,
# enlazando por Nómina/Código contra el Empleado que este comando ya dejó
# creado.
import datetime
from pathlib import Path

import openpyxl
from django.core.exceptions import ValidationError
from django.core.management.base import BaseCommand, CommandError

from apps.employment.models import Empleado
from apps.imports.models import ColaboradorRawRow, ImportBatch, RawValueAlias
from apps.imports.normalize import resolve_against_catalog
from apps.persons.models import EstadoCivil, Persona

REQUIRED_COLUMNS = ["Código", "Nombre", "Apellido Paterno", "Apellido Materno"]

# Campos secundarios: si vienen mal capturados en la fuente (RFC con más
# caracteres de los que debería, correo mal formado...), se descartan solos
# en vez de tronar el alta completa de la Persona — un dato roto en un campo
# opcional no debe impedir registrar a alguien real.
DROPPABLE_FIELDS = {
    "rfc", "nss", "curp", "personal_email", "phone",
    "postal_code", "city", "municipality", "state", "address_line",
}


class Command(BaseCommand):
    help = (
        "Importa 'Lista Colaboradores' de GPA: crea/actualiza Persona + "
        "Empleado. No crea Posicion/Contrato — eso lo hace el import de la "
        "sábana de Posiciones."
    )

    def add_arguments(self, parser):
        parser.add_argument("path")
        parser.add_argument("--sheet", default=None)

    def handle(self, *args, **options):
        path = Path(options["path"])
        if not path.exists():
            raise CommandError(f"No existe el archivo: {path}")

        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise CommandError(f"No se pudo abrir {path}: {exc}")

        sheet_name = options["sheet"] or wb.sheetnames[0]
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        header = [h.strip() if isinstance(h, str) else h for h in rows[0]]
        col = {name: i for i, name in enumerate(header) if name}

        missing = [c for c in REQUIRED_COLUMNS if c not in col]
        if missing:
            raise CommandError(f"Faltan columnas esperadas en el archivo: {missing}")

        batch = ImportBatch.objects.create(
            source=ImportBatch.SOURCE_COLABORADORES,
            original_filename=path.name,
        )

        created_persona = 0
        updated_persona = 0
        created_empleado = 0
        skipped_no_codigo = 0
        errors = []
        unresolved_estado_civil = set()

        for i, row in enumerate(rows[1:], start=2):
            if all(v is None for v in row):
                continue

            data = {header[j]: self._cell_to_str(v) for j, v in enumerate(row) if header[j]}
            ColaboradorRawRow.objects.create(import_batch=batch, row_number=i, data=data)

            codigo = self._get(row, col, "Código")
            if not codigo:
                skipped_no_codigo += 1
                continue
            work_number = str(codigo).strip().upper()

            estado_civil_raw = self._get(row, col, "Estado Civil")
            marital_status = None
            if estado_civil_raw:
                marital_status, normalized = resolve_against_catalog(
                    estado_civil_raw, EstadoCivil.objects.all(), "estado_civil", RawValueAlias,
                )
                if marital_status is None:
                    unresolved_estado_civil.add(normalized)

            persona_fields = dict(
                first_name=self._clean(self._get(row, col, "Nombre")),
                last_name_paternal=self._clean(self._get(row, col, "Apellido Paterno")),
                last_name_maternal=self._clean(self._get(row, col, "Apellido Materno")) or "",
                curp=self._clean(self._get(row, col, "CURP")) or None,
                nss=self._clean(self._get(row, col, "NSS")) or None,
                rfc=self._clean(self._get(row, col, "RFC")) or None,
                birth_date=self._to_date(self._get(row, col, "Fecha de Nacimiento")),
                marital_status=marital_status,
                personal_email=self._clean(self._get(row, col, "Correo electrónico")) or "",
                phone=self._clean(self._get(row, col, "No. Telefono")) or "",
                address_line=self._clean(self._get(row, col, "Domicilio")) or "",
                postal_code=self._clean(self._get(row, col, "Codigo Postal")) or "",
                city=self._clean(self._get(row, col, "CiudadDomicilio")) or "",
                municipality=self._clean(self._get(row, col, "MunicipioDomicilio")) or "",
                state=self._clean(self._get(row, col, "Edo.")) or "",
            )

            empleado = Empleado.objects.filter(work_number=work_number).select_related("persona").first()

            if empleado is not None:
                persona = empleado.persona
                for field, value in persona_fields.items():
                    if value not in (None, ""):
                        setattr(persona, field, value)
                dropped = self._clean_dropping_bad_fields(persona)
                if dropped is None:
                    errors.append(f"Fila {i} ({work_number}, actualización): {self._last_error}")
                    continue
                if dropped:
                    self.stdout.write(self.style.WARNING(f"Fila {i} ({work_number}): descarté {sorted(dropped)} por venir mal capturados."))
                persona.save()
                updated_persona += 1
            else:
                persona = Persona(**persona_fields)
                dropped = self._clean_dropping_bad_fields(persona)
                if dropped is None:
                    errors.append(f"Fila {i} ({work_number}, alta): {self._last_error}")
                    continue
                if dropped:
                    self.stdout.write(self.style.WARNING(f"Fila {i} ({work_number}): descarté {sorted(dropped)} por venir mal capturados."))
                persona.save()
                created_persona += 1

                nuevo_empleado = Empleado(persona=persona, work_number=work_number)
                nuevo_empleado.full_clean()
                nuevo_empleado.save()
                created_empleado += 1

        self.stdout.write(self.style.SUCCESS(
            f"Personas creadas: {created_persona}, actualizadas: {updated_persona}. "
            f"Empleados creados: {created_empleado}. Filas sin Código: {skipped_no_codigo}."
        ))
        if unresolved_estado_civil:
            self.stdout.write(self.style.WARNING(
                f"Estado Civil sin resolver ({len(unresolved_estado_civil)} valores distintos): "
                f"{sorted(unresolved_estado_civil)}"
            ))
        if errors:
            self.stdout.write(self.style.ERROR(f"{len(errors)} filas con error, no se crearon/actualizaron:"))
            for err in errors[:20]:
                self.stdout.write(self.style.ERROR(f"  {err}"))
            if len(errors) > 20:
                self.stdout.write(self.style.ERROR(f"  ... y {len(errors) - 20} más."))

    def _clean_dropping_bad_fields(self, persona):
        """
        Intenta full_clean(); si falla solo por campos "descartables"
        (DROPPABLE_FIELDS), los limpia a None/"" y reintenta una vez. Si
        falla por cualquier otro campo (los realmente obligatorios), regresa
        None y deja el error en self._last_error para que el llamador lo
        reporte. Regresa el set de campos descartados (vacío si no hubo que
        descartar ninguno).
        """
        try:
            persona.full_clean()
            return set()
        except ValidationError as exc:
            failing_fields = set(exc.message_dict.keys())
            if not failing_fields <= DROPPABLE_FIELDS:
                self._last_error = exc
                return None
            for field in failing_fields:
                setattr(persona, field, None if field in {"rfc", "nss", "curp"} else "")
            try:
                persona.full_clean()
            except ValidationError as exc2:
                self._last_error = exc2
                return None
            return failing_fields

    @staticmethod
    def _get(row, col, name):
        idx = col.get(name)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    @staticmethod
    def _clean(value):
        if value is None:
            return None
        text = str(value).strip()
        return text or None

    @staticmethod
    def _cell_to_str(value):
        if value is None:
            return None
        if isinstance(value, (datetime.datetime, datetime.date)):
            return value.isoformat()
        return str(value)

    @staticmethod
    def _to_date(value):
        if isinstance(value, datetime.datetime):
            return value.date()
        if isinstance(value, datetime.date):
            return value
        return None

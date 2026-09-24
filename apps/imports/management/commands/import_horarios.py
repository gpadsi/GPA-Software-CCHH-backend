# apps/imports/management/commands/import_horarios.py
# Importa "Lista relación horarios empleados": siembra TipoHorario (limpio,
# 1 a 1 en la fuente) y crea AsignacionHorario por Empleado directamente —
# no depende de que exista Contrato/Posición todavía (Empleado ya existe
# desde el import de Colaboradores).
#
# `catorcena` se deja sin resolver a propósito (ver apps/schedules/models.py)
# — se guarda `fecha_referencia` (Fec Reg Sis) mientras no exista el
# calendario real de catorcenas de GPA.
import datetime
from pathlib import Path

import openpyxl
from django.core.management.base import BaseCommand, CommandError

from apps.employment.models import Empleado
from apps.imports.models import HorarioRawRow, ImportBatch
from apps.schedules.models import AsignacionHorario, TipoHorario

REQUIRED_COLUMNS = ["NoEmpleado", "TipoHorario", "Horario", "Fec Reg Sis"]


class Command(BaseCommand):
    help = (
        "Importa 'Lista relación horarios empleados': siembra TipoHorario y "
        "crea AsignacionHorario por Empleado."
    )

    def add_arguments(self, parser):
        parser.add_argument("path")
        parser.add_argument("--sheet", default=None)

    def handle(self, *args, **options):
        path = Path(options["path"])
        if not path.exists():
            raise CommandError(f"No existe el archivo: {path}")

        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as exc:
            raise CommandError(f"No se pudo abrir {path}: {exc}")

        sheet_name = options["sheet"] or wb.sheetnames[0]
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        header = [h.strip() if isinstance(h, str) else h for h in rows[0]]
        col = {name: i for i, name in enumerate(header) if name}

        missing = [c for c in REQUIRED_COLUMNS if c not in col]
        if missing:
            raise CommandError(f"Faltan columnas esperadas en el archivo: {missing}")

        batch = ImportBatch.objects.create(
            source=ImportBatch.SOURCE_HORARIOS,
            original_filename=path.name,
        )

        tipos_creados = 0
        asignaciones_creadas = 0
        sin_nomina = 0
        empleado_no_encontrado = 0
        sin_tipo_horario = 0

        for i, row in enumerate(rows[1:], start=2):
            if all(v is None for v in row):
                continue

            data = {header[j]: self._cell_to_str(v) for j, v in enumerate(row) if header[j]}
            HorarioRawRow.objects.create(import_batch=batch, row_number=i, data=data)

            no_empleado = self._get(row, col, "NoEmpleado")
            tipo_codigo = self._get(row, col, "TipoHorario")
            horario_texto = self._get(row, col, "Horario")
            fec_reg = self._get(row, col, "Fec Reg Sis")

            tipo_horario = None
            if tipo_codigo:
                tipo_horario, was_created = TipoHorario.objects.get_or_create(
                    name=str(tipo_codigo).strip(),
                    defaults={"descripcion": str(horario_texto).strip() if horario_texto else ""},
                )
                tipos_creados += int(was_created)

            if not no_empleado:
                sin_nomina += 1
                continue

            work_number = str(no_empleado).strip().upper()
            empleado = Empleado.objects.filter(work_number=work_number).first()
            if empleado is None:
                empleado_no_encontrado += 1
                continue

            if tipo_horario is None:
                sin_tipo_horario += 1
                continue

            fecha_referencia = self._to_date(fec_reg) or datetime.date.today()

            _, was_created = AsignacionHorario.objects.get_or_create(
                empleado=empleado,
                fecha_referencia=fecha_referencia,
                defaults={"tipo_horario": tipo_horario},
            )
            asignaciones_creadas += int(was_created)

        self.stdout.write(self.style.SUCCESS(
            f"TipoHorario creados: {tipos_creados}. AsignacionHorario creadas: {asignaciones_creadas}."
        ))
        self.stdout.write(self.style.WARNING(
            f"Sin Nómina (no vinculadas): {sin_nomina}. Empleado no encontrado: {empleado_no_encontrado}. "
            f"Sin TipoHorario en la fila: {sin_tipo_horario}."
        ))

    @staticmethod
    def _get(row, col, name):
        idx = col.get(name)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    @staticmethod
    def _cell_to_str(value):
        if value is None:
            return None
        if isinstance(value, (datetime.datetime, datetime.date)):
            return value.isoformat()
        return str(value)

    @staticmethod
    def _to_date(value):
        if isinstance(value, datetime.datetime):
            return value.date()
        if isinstance(value, datetime.date):
            return value
        return None

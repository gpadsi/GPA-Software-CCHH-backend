import tempfile
from datetime import date
from io import StringIO
from pathlib import Path

import openpyxl
from django.contrib.contenttypes.models import ContentType
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase

from apps.employment.models import Contrato, Empleado
from apps.imports.models import ImportBatch, RawValueAlias
from apps.imports.normalize import normalize_text, resolve_against_catalog
from apps.locations.models import Area
from apps.persons.models import Escolaridad, EstadoCivil, Genero, Persona
from apps.positions.models import Posicion


class NormalizeTextTests(TestCase):
    def test_none_becomes_empty_string(self):
        self.assertEqual(normalize_text(None), "")

    def test_uppercases_and_strips_extra_whitespace(self):
        self.assertEqual(normalize_text("  matriz  "), "MATRIZ")

    def test_strips_accents(self):
        self.assertEqual(normalize_text("Ingeniería"), "INGENIERIA")

    def test_punctuation_becomes_whitespace_and_collapses(self):
        self.assertEqual(normalize_text("ING."), "ING")
        self.assertEqual(normalize_text("Auxiliar, de Producción"), "AUXILIAR DE PRODUCCION")

    def test_inserts_space_between_letters_and_digits_even_without_domain(self):
        self.assertEqual(normalize_text("N1"), "N 1")

    def test_domain_synonym_substitution_applies_per_token(self):
        self.assertEqual(normalize_text("AUX", domain="puesto"), "AUXILIAR")
        self.assertEqual(normalize_text("N1", domain="nave"), "NAVE 1")

    def test_synonyms_are_scoped_to_their_own_domain(self):
        # "N" solo se expande a "NAVE" en el dominio "nave" — en "puesto" no
        # existe ese sinónimo, así que se queda tal cual.
        self.assertEqual(normalize_text("N1", domain="puesto"), "N 1")

    def test_domain_word_merge_replaces_the_whole_normalized_phrase(self):
        self.assertEqual(normalize_text("Operador de Soldador", domain="puesto"), "OPERADOR DE SOLDADURA")

    def test_unmapped_value_in_a_domain_is_returned_unchanged(self):
        self.assertEqual(normalize_text("Valor Sin Mapear", domain="puesto"), "VALOR SIN MAPEAR")


class ResolveAgainstCatalogTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.masculino = Genero.objects.create(name="Masculino")

    def test_resolves_by_exact_normalized_name_match(self):
        obj, normalized = resolve_against_catalog("  masculino  ", Genero.objects.all(), "genero", RawValueAlias)
        self.assertEqual(obj, self.masculino)
        self.assertEqual(normalized, "MASCULINO")

    def test_returns_none_and_the_normalized_value_when_nothing_matches(self):
        obj, normalized = resolve_against_catalog("Otro Valor", Genero.objects.all(), "genero", RawValueAlias)
        self.assertIsNone(obj)
        self.assertEqual(normalized, "OTRO VALOR")

    def test_never_creates_anything_when_unresolved(self):
        resolve_against_catalog("Valor Inventado", Genero.objects.all(), "genero", RawValueAlias)
        self.assertFalse(Genero.objects.filter(name="Valor Inventado").exists())

    def test_a_confirmed_alias_resolves_even_when_the_text_does_not_match_the_name(self):
        content_type = ContentType.objects.get_for_model(Genero)
        RawValueAlias.objects.create(
            domain="genero", raw_value="MASC", raw_value_original="Masc",
            content_type=content_type, object_id=str(self.masculino.pk),
        )

        obj, normalized = resolve_against_catalog("Masc", Genero.objects.all(), "genero", RawValueAlias)
        self.assertEqual(obj, self.masculino)
        self.assertEqual(normalized, "MASC")


class ImportBatchChecksumTests(TestCase):
    @staticmethod
    def _write_horarios_xlsx(path, no_empleado="ADV0001"):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["NoEmpleado", "TipoHorario", "Horario", "Fec Reg Sis"])
        ws.append([no_empleado, "H01", "07:00 - 16:00", "2026-01-01"])
        wb.save(path)

    def test_checksum_is_stable_across_filenames_and_sensitive_to_content(self):
        with tempfile.TemporaryDirectory() as tmp:
            same_content_a = Path(tmp) / "a.xlsx"
            same_content_b = Path(tmp) / "b.xlsx"  # mismo contenido, nombre distinto
            different_content = Path(tmp) / "c.xlsx"
            self._write_horarios_xlsx(same_content_a)
            self._write_horarios_xlsx(same_content_b)
            self._write_horarios_xlsx(different_content, no_empleado="ADV9999")

            self.assertEqual(
                ImportBatch.checksum_for(same_content_a), ImportBatch.checksum_for(same_content_b)
            )
            self.assertNotEqual(
                ImportBatch.checksum_for(same_content_a), ImportBatch.checksum_for(different_content)
            )

    def test_find_duplicate_matches_by_source_and_checksum_only(self):
        batch = ImportBatch.objects.create(
            source=ImportBatch.SOURCE_HORARIOS, original_filename="horarios.xlsx", file_checksum="abc123",
        )
        self.assertEqual(ImportBatch.find_duplicate(ImportBatch.SOURCE_HORARIOS, "abc123"), batch)
        self.assertIsNone(ImportBatch.find_duplicate(ImportBatch.SOURCE_COLABORADORES, "abc123"))
        self.assertIsNone(ImportBatch.find_duplicate(ImportBatch.SOURCE_HORARIOS, "otro-checksum"))

    def test_find_duplicate_ignores_blank_checksums_from_lotes_previos_al_campo(self):
        ImportBatch.objects.create(source=ImportBatch.SOURCE_HORARIOS, original_filename="viejo.xlsx")
        self.assertIsNone(ImportBatch.find_duplicate(ImportBatch.SOURCE_HORARIOS, ""))


class ImportHorariosDuplicateGuardTests(TestCase):
    """
    Prueba de extremo a extremo (comando real, archivo .xlsx real) del caso
    que ya pasó de verdad: horarios.xlsx importado dos veces sin querer
    (ver docs/DATA_GAPS_JUSTIFICATION.md). import_horarios es el más simple
    de los tres comandos que usan este mismo guardián (import_colaboradores
    e import_posiciones comparten exactamente el mismo patrón).
    """
    @staticmethod
    def _write_horarios_xlsx(path, no_empleado="ADV0001"):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["NoEmpleado", "TipoHorario", "Horario", "Fec Reg Sis"])
        ws.append([no_empleado, "H01", "07:00 - 16:00", "2026-01-01"])
        wb.save(path)

    def test_reimportar_el_mismo_archivo_se_bloquea_sin_force(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "horarios.xlsx"
            self._write_horarios_xlsx(path)

            call_command("import_horarios", str(path), stdout=StringIO())
            self.assertEqual(ImportBatch.objects.count(), 1)

            with self.assertRaises(CommandError):
                call_command("import_horarios", str(path), stdout=StringIO())
            self.assertEqual(ImportBatch.objects.count(), 1)

    def test_force_permite_reimportar_el_mismo_archivo(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "horarios.xlsx"
            self._write_horarios_xlsx(path)

            call_command("import_horarios", str(path), stdout=StringIO())
            call_command("import_horarios", str(path), force=True, stdout=StringIO())
            self.assertEqual(ImportBatch.objects.count(), 2)

    def test_mismo_nombre_con_contenido_corregido_no_se_bloquea(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "horarios.xlsx"
            self._write_horarios_xlsx(path, no_empleado="ADV0001")
            call_command("import_horarios", str(path), stdout=StringIO())

            self._write_horarios_xlsx(path, no_empleado="ADV0002")  # sábana corregida, mismo nombre
            call_command("import_horarios", str(path), stdout=StringIO())
            self.assertEqual(ImportBatch.objects.count(), 2)


class ImportPosicionesEmpleadoYPersonaTests(TestCase):
    """
    import_posiciones con un .xlsx real: cómo encuentra al Empleado de cada
    fila (Nómina confirmada por nombre, nombre exacto cuando no hay Nómina,
    nunca entre homónimos), qué rellena de la Persona y que un Área del
    catálogo no se pierda por llamarse igual que una unidad.
    """
    ENCABEZADOS = [
        "Nomina", "Empresa", "Unidad de Negocio ", "Área", "Nombre", "Genero ", "Escolaridad",
        "Estado Civil", "Es padre", "Estatus", "Fecha de Ingreso como Colaborador",
        "Fecha de Baja como Colaborador",
    ]

    @classmethod
    def setUpTestData(cls):
        for seed in (
            "seed_tenant", "seed_organizational_levels", "seed_gpa_companies",
            "seed_position_catalogs", "seed_persons_catalogs", "seed_baja_catalogs",
        ):
            call_command(seed, stdout=StringIO())
        Escolaridad.objects.create(name="Secundaria")
        EstadoCivil.objects.create(name="Casado")
        cls.area_cei = Area.objects.create(code="CEI", name="Cei")

    @staticmethod
    def _empleado(work_number, paterno, materno, nombre):
        persona = Persona.objects.create(first_name=nombre, last_name_paternal=paterno, last_name_maternal=materno)
        return Empleado.objects.create(persona=persona, work_number=work_number)

    def _importar(self, *filas):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "sabana.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Posiciones"
            for j, encabezado in enumerate(self.ENCABEZADOS, start=1):
                ws.cell(row=4, column=j, value=encabezado)
            for i, fila in enumerate(filas, start=5):
                for j, encabezado in enumerate(self.ENCABEZADOS, start=1):
                    ws.cell(row=i, column=j, value=fila.get(encabezado.strip()))
            wb.save(path)
            call_command("import_posiciones", str(path), stdout=StringIO())

    def test_fila_sin_nomina_se_liga_por_nombre_exacto_y_rellena_la_persona(self):
        empleado = self._empleado("ADV9001", "PEREZ", "LOPEZ", "JUAN")
        self._importar({
            "Empresa": "GPA ADVANCED MANUFACTURING", "Unidad de Negocio": "Tooling",
            "Nombre": "PÉREZ LÓPEZ JUAN", "Genero": "MASCULINO", "Escolaridad": "Secundaria",
            "Estado Civil": "Casado", "Es padre": "Si", "Estatus": "COLABORADOR ACTIVO",
            "Fecha de Ingreso como Colaborador": date(2026, 8, 24),
        })

        contrato = Contrato.objects.get(empleado=empleado)
        self.assertEqual(contrato.fecha_ingreso, date(2026, 8, 24))
        self.assertEqual(contrato.posicion.organization_node.name, "Tooling")
        persona = Persona.objects.get(pk=empleado.persona_id)
        self.assertEqual(persona.education_level.name, "Secundaria")
        self.assertEqual(persona.marital_status.name, "Casado")
        self.assertEqual(persona.gender.name, "Masculino")
        self.assertTrue(persona.has_children)

    def test_nomina_confirmada_por_nombre_crea_el_contrato_del_empleado_real(self):
        empleado = self._empleado("ADV1072", "CONTRERAS", "TRUJILLO", "MIGUEL ANGEL")
        self._importar({
            "Nomina": "ADD0170", "Empresa": "GPA ADDITRONICS", "Nombre": "CONTRERAS TRUJILLO MIGUEL ANGEL",
            "Estatus": "COLABORADOR BAJA", "Fecha de Ingreso como Colaborador": date(2025, 1, 6),
            "Fecha de Baja como Colaborador": date(2026, 8, 20),
        })

        contrato = Contrato.objects.get(empleado=empleado)
        self.assertEqual(contrato.fecha_baja, date(2026, 8, 20))

    def test_nomina_confirmada_no_se_usa_si_el_nombre_ya_no_coincide(self):
        self._empleado("ADV1072", "CONTRERAS", "TRUJILLO", "MIGUEL ANGEL")
        self._importar({
            "Nomina": "ADD0170", "Empresa": "GPA ADDITRONICS", "Nombre": "OTRA PERSONA DISTINTA",
            "Estatus": "COLABORADOR BAJA", "Fecha de Ingreso como Colaborador": date(2025, 1, 6),
        })

        self.assertEqual(Posicion.objects.count(), 1)
        self.assertFalse(Contrato.objects.exists())

    def test_homonimos_no_se_ligan_aunque_solo_uno_sea_empleado(self):
        self._empleado("ADV9002", "GOMEZ", "RUIZ", "ANA")
        Persona.objects.create(first_name="ANA", last_name_paternal="GOMEZ", last_name_maternal="RUIZ")
        self._importar({
            "Empresa": "GPA ADVANCED MANUFACTURING", "Nombre": "GOMEZ RUIZ ANA",
            "Estatus": "COLABORADOR ACTIVO", "Fecha de Ingreso como Colaborador": date(2026, 1, 5),
        })

        self.assertFalse(Contrato.objects.exists())

    def test_area_del_catalogo_no_se_descarta_por_llamarse_igual_que_una_unidad(self):
        # «CEI» es la unidad de CEI Aerospace Group (seed_gpa_companies) y
        # también un Área ya confirmada en el catálogo.
        self._importar({
            "Empresa": "CEI AEROSPACE GROUP", "Área": "CEI", "Nombre": "VACP", "Estatus": "VACANTE ACTIVA",
        })

        self.assertEqual(Posicion.objects.get().area, self.area_cei)
